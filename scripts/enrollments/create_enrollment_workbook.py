import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADERS = ["phone_number", "course", "active", "student_name", "notes"]
COURSES = ["CMA", "CPA", "CFA", "ACCA", "CS", "EA"]


def create_workbook(
    path: Path,
    *,
    with_sample: bool = False,
    force: bool = False,
) -> None:
    if path.exists() and not force:
        raise FileExistsError(f"Refusing to overwrite existing workbook: {path}")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Enrollments"

    sheet.append(HEADERS)
    if with_sample:
        sheet.append(["919999999999", "CMA", "YES", "Sample student", "Example only"])

    header_fill = PatternFill("solid", fgColor="17365D")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:E10000"
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 12
    sheet.column_dimensions["D"].width = 24
    sheet.column_dimensions["E"].width = 42

    if with_sample:
        sheet.cell(row=2, column=1).number_format = "@"

    course_validation = DataValidation(
        type="list",
        formula1='"CMA,CPA,CFA,ACCA,CS,EA"',
        allow_blank=False,
        error="Select CMA, CPA, CFA, ACCA, CS, or EA.",
        errorTitle="Invalid course",
    )
    active_validation = DataValidation(
        type="list",
        formula1='"YES,NO"',
        allow_blank=False,
        error="Select YES or NO.",
        errorTitle="Invalid access status",
    )
    sheet.add_data_validation(course_validation)
    sheet.add_data_validation(active_validation)
    course_validation.add("B2:B10000")
    active_validation.add("C2:C10000")

    disabled_fill = PatternFill("solid", fgColor="F4CCCC")
    sheet.conditional_formatting.add(
        "A2:E10000",
        FormulaRule(formula=['$C2="NO"'], fill=disabled_fill),
    )

    table = Table(displayName="WhatsAppEnrollments", ref=f"A1:E{max(2, sheet.max_row)}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)

    instructions = workbook.create_sheet("Instructions")
    instructions.column_dimensions["A"].width = 110
    instructions["A1"] = "NorthStar Academy WhatsApp Access Control"
    instructions["A1"].font = Font(size=16, bold=True, color="17365D")
    instructions["A3"] = "1. Open the Enrollments sheet."
    instructions["A4"] = "2. Add one student per row. Enter a 10-digit Indian number or a full country-code number."
    instructions["A5"] = "3. Add one row per student/course. Repeat the phone number to grant multiple courses."
    instructions["A6"] = "4. Select CMA, CPA, CFA, ACCA, CS, or EA, then set active to YES or NO."
    instructions["A7"] = "5. Save this same file. The running bot and dashboard reload changes automatically."
    instructions["A9"] = "Important: Keep each phone/course pair unique. For duplicate pairs, the last row wins."
    instructions["A10"] = "Students see only their active enrolled courses and can switch between those courses in WhatsApp."
    instructions["A11"] = "Rows with invalid phone numbers, courses, or active values are denied access."

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Create the WhatsApp enrollment Excel workbook.")
    parser.add_argument(
        "output",
        nargs="?",
        default="/app/data/whatsapp_enrollments.xlsx",
        help="Output .xlsx path.",
    )
    parser.add_argument(
        "--with-sample",
        action="store_true",
        help="Add one clearly marked example enrollment. The default workbook is empty.",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Replace an existing workbook. Without this flag existing data is preserved.",
    )
    args = parser.parse_args()
    output = Path(args.output)
    try:
        create_workbook(output, with_sample=args.with_sample, force=args.force)
    except FileExistsError as exc:
        parser.error(str(exc))
    print(f"Created enrollment workbook: {output}")


if __name__ == "__main__":
    main()

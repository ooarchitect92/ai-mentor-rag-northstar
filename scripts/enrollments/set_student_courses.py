import argparse
import os
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import load_workbook


COURSES = ["CMA", "CPA", "CFA", "ACCA", "CS", "EA"]


def normalize_phone(value: object) -> str:
    text = "" if value is None else str(value)
    digits = "".join(char for char in text if char.isdigit())
    if not 8 <= len(digits) <= 15:
        raise ValueError("Phone number must contain 8 to 15 digits, including country code")
    return digits


def set_student_courses(path: Path, phone: str, courses: list[str], student_name: str = "") -> None:
    workbook = load_workbook(path)
    sheet = workbook["Enrollments"]
    normalized_phone = normalize_phone(phone)
    selected = [course.upper() for course in courses]
    invalid = sorted(set(selected) - set(COURSES))
    if invalid:
        raise ValueError(f"Unsupported courses: {', '.join(invalid)}")

    existing_rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        if normalize_phone(row[0]) != normalized_phone:
            existing_rows.append(list(row[:5]))

    sheet.delete_rows(2, max(1, sheet.max_row - 1))
    for course in COURSES:
        if course in selected:
            sheet.append([normalized_phone, course, "YES", student_name, "Multi-course enrollment"])
    for row in existing_rows:
        sheet.append(row)

    for validation in sheet.data_validations.dataValidation:
        if "B" in str(validation.sqref):
            validation.formula1 = '"CMA,CPA,CFA,ACCA,CS,EA"'
            validation.error = "Select CMA, CPA, CFA, ACCA, CS, or EA."

    if "WhatsAppEnrollments" in sheet.tables:
        sheet.tables["WhatsAppEnrollments"].ref = f"A1:E{max(2, sheet.max_row)}"

    instructions = workbook["Instructions"]
    instructions["A5"] = "3. Add one row per student/course. Repeat the phone number to grant multiple courses."
    instructions["A6"] = "4. Select CMA, CPA, CFA, ACCA, CS, or EA, then set active to YES or NO."
    instructions["A9"] = "Important: Keep each phone/course pair unique. For duplicate pairs, the last row wins."
    instructions["A10"] = "Students see only their active enrolled courses and can switch between those courses in WhatsApp."

    temporary_path: Path | None = None
    try:
        with NamedTemporaryFile(
            dir=path.parent,
            prefix=f".{path.name}.",
            suffix=".tmp.xlsx",
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
        workbook.save(temporary_path)
        workbook.close()
        with temporary_path.open("r+b") as saved:
            saved.flush()
            os.fsync(saved.fileno())
        os.replace(temporary_path, path)
        temporary_path = None
    finally:
        workbook.close()
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser(description="Set all active courses for one WhatsApp student.")
    parser.add_argument("phone")
    parser.add_argument("courses", nargs="+", choices=COURSES)
    parser.add_argument("--name", default="")
    parser.add_argument("--workbook", default="/app/data/whatsapp_enrollments.xlsx")
    args = parser.parse_args()
    set_student_courses(Path(args.workbook), args.phone, args.courses, args.name)
    print(f"Updated {args.phone}: {', '.join(args.courses)}")


if __name__ == "__main__":
    main()

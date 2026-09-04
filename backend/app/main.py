import asyncio
import ipaddress
import json
import logging
import secrets
import time
from contextlib import asynccontextmanager
from io import BytesIO
from pathlib import Path
from urllib.parse import quote, urlparse
from fastapi import Depends, FastAPI, File, Form, HTTPException, Query, Request, UploadFile, status
from fastapi.responses import FileResponse, ORJSONResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from .admin_api import create_knowledge_documents, router as admin_router
from .admin_store import AdminStore
from .auth import require_admin_token
from .cache import Cache
from .config import get_settings
from .mentor import MentorService
from .schemas import (
    ChatRequest,
    ChatResponse,
    Course,
    DocumentType,
    IngestResult,
    WhatsAppEnrollmentRequest,
    WhatsAppEnrollmentResponse,
    WhatsAppMenuRequest,
    WhatsAppStartRequest,
)
from .training import cancel_training_tasks, run_training_job, training_supervisor
from .vector_store import VectorStore
from .webhook_queue import (
    cancel_whatsapp_webhook_tasks,
    schedule_whatsapp_webhook,
    whatsapp_webhook_supervisor,
)
from .whatsapp import (
    COURSES,
    EnrollmentSourceReadOnlyError,
    EnrollmentWorkbookError,
    WhatsAppClient,
    enrollment_source,
    enrollment_roster,
    bulk_set_enrollment_states,
    get_enrolled_courses,
    normalize_wa_id,
    ordered_courses,
    remove_enrolled_course,
    set_enrolled_course,
    split_feedback_marker,
    verify_meta_signature,
)


logger = logging.getLogger("uvicorn.error")

_DIRECT_WHATSAPP_CALLBACK_PATHS = {
    "/v1/whatsapp/ziplin/webhook",
}
_WHATSAPP_WEBHOOK_MAX_BODY_BYTES = 2 * 1024 * 1024


def _direct_whatsapp_callback_state(callback_url: str) -> tuple[bool, bool]:
    """Return (direct_path, stable_https_origin) for the configured Meta callback."""
    if not callback_url:
        return False, False
    try:
        parsed = urlparse(callback_url.strip())
        hostname = (parsed.hostname or "").rstrip(".").casefold()
    except ValueError:
        return False, False
    callback_path = parsed.path
    is_direct = bool(
        callback_path in _DIRECT_WHATSAPP_CALLBACK_PATHS
        and not parsed.params
        and not parsed.query
        and not parsed.fragment
    )
    public_host = bool(hostname and "." in hostname)
    try:
        ipaddress.ip_address(hostname)
    except ValueError:
        public_host = public_host and hostname not in {"localhost"} and not hostname.endswith(
            (".localhost", ".local")
        )
    else:
        # A permanent DNS hostname is required even when a literal IP address
        # is publicly routable; certificates, ownership and rotation are much
        # safer to operate against a stable name.
        public_host = False
    try:
        supported_port = parsed.port in {None, 443}
    except ValueError:
        supported_port = False
    is_stable_https = bool(
        parsed.scheme.casefold() == "https"
        and public_host
        and supported_port
        and not parsed.username
        and not parsed.password
        and not hostname.endswith(".trycloudflare.com")
    )
    return is_direct, is_stable_https


def validate_startup_configuration() -> None:
    settings = get_settings()
    if settings.app_environment.casefold() != "production":
        return
    errors: list[str] = []
    admin_token = settings.admin_token.strip()
    if len(admin_token) < 24 or admin_token == "change-me":
        errors.append("ADMIN_TOKEN must be a unique value of at least 24 characters")
    verify_token = settings.whatsapp_verify_token.strip()
    if not verify_token or verify_token == "change-me-whatsapp":
        errors.append("WHATSAPP_VERIFY_TOKEN is required and must not use the example value")
    app_secret = settings.whatsapp_app_secret.strip() or settings.meta_app_secret.strip()
    if not app_secret:
        errors.append("WHATSAPP_APP_SECRET is required for direct Meta webhook signature verification")
    callback_url = settings.whatsapp_webhook_callback_url.strip()
    direct_callback, stable_callback = _direct_whatsapp_callback_state(callback_url)
    if not callback_url:
        errors.append("WHATSAPP_WEBHOOK_CALLBACK_URL is required in production")
    elif not direct_callback:
        errors.append(
            "WHATSAPP_WEBHOOK_CALLBACK_URL must point directly to /v1/whatsapp/ziplin/webhook"
        )
    elif not stable_callback:
        errors.append("WHATSAPP_WEBHOOK_CALLBACK_URL must use a permanent HTTPS hostname")
    if not (settings.whatsapp_access_token.strip() or settings.whatsapp_token.strip()):
        errors.append("WHATSAPP_ACCESS_TOKEN is required")
    if not settings.whatsapp_phone_number_id.strip():
        errors.append("WHATSAPP_PHONE_NUMBER_ID is required")
    if not settings.whatsapp_business_account_id.strip():
        errors.append("WHATSAPP_BUSINESS_ACCOUNT_ID is required for webhook WABA isolation")
    feedback_digits = "".join(character for character in settings.whatsapp_feedback_number if character.isdigit())
    if not 8 <= len(feedback_digits) <= 15:
        errors.append("WHATSAPP_FEEDBACK_NUMBER must contain 8 to 15 E.164 digits")
    feedback_marker, _ = split_feedback_marker(settings.whatsapp_feedback_prefill)
    if not feedback_marker:
        errors.append("WHATSAPP_FEEDBACK_PREFILL must begin with the FEEDBACK command marker")
    provider_keys = {
        "nvidia": settings.nvidia_api_key.strip(),
        "gemini": settings.gemini_api_key.strip(),
        "anthropic": settings.anthropic_api_key.strip(),
    }
    if not provider_keys.get(settings.mentor_provider):
        errors.append(f"the configured {settings.mentor_provider.upper()} mentor provider key is required")
    if settings.whatsapp_messaging_enabled and not provider_keys["gemini"]:
        errors.append("GEMINI_API_KEY is required for WhatsApp image questions")
    if settings.mentor_policy_provider == "gemini" and not provider_keys["gemini"]:
        errors.append("GEMINI_API_KEY is required by MENTOR_POLICY_PROVIDER=gemini")
    fallback_provider = settings.mentor_fallback_provider
    if fallback_provider != "none" and not provider_keys.get(fallback_provider):
        errors.append(f"the configured {fallback_provider.upper()} fallback provider key is required")
    if settings.embedding_provider == "openai" and not (settings.openai_api_key or "").strip():
        errors.append("OPENAI_API_KEY is required by EMBEDDING_PROVIDER=openai")
    minimum_lease = max(60, int(settings.request_timeout_seconds) * 4)
    if settings.training_lease_timeout_seconds < minimum_lease:
        errors.append(
            f"TRAINING_LEASE_TIMEOUT_SECONDS must be at least {minimum_lease} for configured request timeouts"
        )
    if not 5 <= settings.training_recovery_interval_seconds < settings.training_lease_timeout_seconds:
        errors.append(
            "TRAINING_RECOVERY_INTERVAL_SECONDS must be at least 5 and shorter than the lease timeout"
        )
    minimum_webhook_lease = (
        settings.whatsapp_inbound_processing_ttl_seconds
        + max(5, settings.whatsapp_webhook_recovery_interval_seconds)
    )
    if settings.whatsapp_webhook_lease_timeout_seconds < minimum_webhook_lease:
        errors.append(
            "WHATSAPP_WEBHOOK_LEASE_TIMEOUT_SECONDS must be at least "
            f"{minimum_webhook_lease} so in-flight message deduplication expires before recovery"
        )
    if not 5 <= settings.whatsapp_webhook_recovery_interval_seconds < settings.whatsapp_webhook_lease_timeout_seconds:
        errors.append(
            "WHATSAPP_WEBHOOK_RECOVERY_INTERVAL_SECONDS must be at least 5 and shorter than the lease timeout"
        )
    if not 1 <= settings.whatsapp_webhook_max_attempts <= 20:
        errors.append("WHATSAPP_WEBHOOK_MAX_ATTEMPTS must be between 1 and 20")
    if not 30 <= settings.whatsapp_inbound_processing_ttl_seconds <= 3600:
        errors.append("WHATSAPP_INBOUND_PROCESSING_TTL_SECONDS must be between 30 and 3600")
    if errors:
        raise RuntimeError("Unsafe production configuration: " + "; ".join(errors))


@asynccontextmanager
async def lifespan(_: FastAPI):
    validate_startup_configuration()
    supervisors = [
        asyncio.create_task(training_supervisor(), name="knowledge-training-supervisor"),
        asyncio.create_task(whatsapp_webhook_supervisor(), name="whatsapp-webhook-supervisor"),
    ]
    try:
        yield
    finally:
        for supervisor in supervisors:
            supervisor.cancel()
        await asyncio.gather(*supervisors, return_exceptions=True)
        await cancel_training_tasks()
        await cancel_whatsapp_webhook_tasks()


app = FastAPI(
    title="AI Mentor RAG API",
    description="AI mentor for CMA / CPA / CFA / ACCA / CS / EA training institutes",
    version="1.1.0",
    default_response_class=ORJSONResponse,
    lifespan=lifespan,
)

project_dir = Path(__file__).resolve().parents[2]
frontend_dir = project_dir / "frontend"
admin_dashboard_dir = project_dir / "admin-dashboard"
if frontend_dir.exists():
    app.mount("/static", StaticFiles(directory=str(frontend_dir)), name="static")
if admin_dashboard_dir.exists():
    app.mount(
        "/admin",
        StaticFiles(directory=str(admin_dashboard_dir), html=True),
        name="admin-dashboard",
    )

app.include_router(admin_router)


@app.middleware("http")
async def add_security_headers(request: Request, call_next):
    response = await call_next(request)
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("X-Frame-Options", "DENY")
    response.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    response.headers.setdefault("Permissions-Policy", "camera=(), microphone=(), geolocation=()")
    if request.url.path.startswith("/admin"):
        response.headers.setdefault(
            "Content-Security-Policy",
            "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; "
            "img-src 'self' data: blob:; connect-src 'self'; object-src 'none'; "
            "base-uri 'none'; frame-ancestors 'none'; form-action 'self'",
        )
    if request.url.path.startswith("/v1/admin"):
        response.headers.setdefault("Cache-Control", "no-store")
    return response


@app.get("/")
async def home():
    index = frontend_dir / "index.html"
    if index.exists():
        return FileResponse(index)
    return {"message": "AI Mentor RAG API is running. Open /docs for API docs."}


@app.get("/health")
async def health():
    settings = get_settings()
    provider_models = {
        "nvidia": settings.nvidia_model,
        "gemini": settings.gemini_model,
        "anthropic": settings.anthropic_model,
    }
    primary_model = provider_models[settings.mentor_provider]
    fallback_model = None
    fallback_configured = False
    if settings.mentor_fallback_provider == "nvidia":
        fallback_model = settings.nvidia_model
        fallback_configured = bool(settings.nvidia_api_key)
    elif settings.mentor_fallback_provider == "gemini":
        fallback_model = settings.gemini_model
        fallback_configured = bool(settings.gemini_api_key)
    elif settings.mentor_fallback_provider == "anthropic":
        fallback_model = settings.anthropic_model
        fallback_configured = bool(settings.anthropic_api_key)

    return {
        "status": "ok",
        "collection": settings.qdrant_collection,
        "llm_provider": settings.mentor_provider,
        "llm_model": primary_model,
        "llm_fallback_provider": settings.mentor_fallback_provider,
        "llm_fallback_model": fallback_model,
        "llm_fallback_configured": fallback_configured,
        "llm_policy_provider": settings.mentor_policy_provider,
        "question_answer_file": bool(settings.question_answer_file),
        "embedding_provider": settings.embedding_provider,
        "whatsapp_configured": bool(
            settings.whatsapp_use_mock
            or ((settings.whatsapp_access_token or settings.whatsapp_token) and settings.whatsapp_phone_number_id)
        ),
        "whatsapp_mock": settings.whatsapp_use_mock,
    }


@app.get("/live")
async def live():
    return {"status": "ok"}


@app.get("/ready")
async def ready():
    async def redis_ready() -> None:
        cache = Cache()
        try:
            if not await asyncio.wait_for(cache.redis.ping(), timeout=3):
                raise RuntimeError("Redis ping failed")
        finally:
            await cache.redis.aclose()

    async def qdrant_ready() -> None:
        vector_store = VectorStore()
        try:
            await asyncio.wait_for(vector_store.ensure_collection(), timeout=3)
        finally:
            await vector_store.client.close()

    async def sqlite_ready() -> None:
        await asyncio.wait_for(
            asyncio.to_thread(lambda: AdminStore()._overview_sync()),
            timeout=3,
        )

    results = await asyncio.gather(
        redis_ready(),
        qdrant_ready(),
        sqlite_ready(),
        return_exceptions=True,
    )
    failures = [
        name
        for name, result in zip(("redis", "qdrant", "sqlite"), results)
        if isinstance(result, Exception)
    ]
    if failures:
        raise HTTPException(status_code=503, detail={"status": "not_ready", "dependencies": failures})
    return {
        "status": "ready",
        "dependencies": {"redis": "ok", "qdrant": "ok", "sqlite": "ok"},
    }


@app.get("/v1/public/config")
async def public_config():
    settings = get_settings()
    digits = "".join(character for character in settings.whatsapp_feedback_number if character.isdigit())
    feedback_url = ""
    callback_url = settings.whatsapp_webhook_callback_url.strip()
    direct_callback, stable_callback = _direct_whatsapp_callback_state(callback_url)
    inbound_delivery_ready = bool(
        direct_callback
        and stable_callback
        and (settings.whatsapp_app_secret or settings.meta_app_secret)
        and settings.whatsapp_verify_token
    )
    feedback_delivery_ready = bool(
        settings.whatsapp_messaging_enabled
        and digits
        and (settings.whatsapp_access_token or settings.whatsapp_token)
        and settings.whatsapp_phone_number_id
        and inbound_delivery_ready
    )
    if feedback_delivery_ready:
        feedback_url = f"https://wa.me/{digits}?text={quote(settings.whatsapp_feedback_prefill)}"
    media_limit_mb = max(1, settings.whatsapp_max_media_bytes // (1024 * 1024))
    return {
        "feedback": {
            "enabled": bool(feedback_url),
            "whatsapp_url": feedback_url,
            "attachment_note": (
                "Screenshots must be attached manually in WhatsApp "
                f"(JPG or PNG, up to {media_limit_mb} MB)."
            ),
        }
    }


@app.get("/v1/whatsapp/ziplin/webhook")
@app.get("/v1/whatsapp/webhook")
async def verify_whatsapp_webhook(
    hub_mode: str | None = Query(default=None, alias="hub.mode"),
    hub_verify_token: str | None = Query(default=None, alias="hub.verify_token"),
    hub_challenge: str | None = Query(default=None, alias="hub.challenge"),
):
    settings = get_settings()
    configured_token = settings.whatsapp_verify_token
    supplied_token = hub_verify_token or ""
    if (
        hub_mode == "subscribe"
        and configured_token.strip()
        and supplied_token.strip()
        and secrets.compare_digest(supplied_token, configured_token)
    ):
        return Response(content=hub_challenge or "", media_type="text/plain")

    raise HTTPException(
        status_code=status.HTTP_403_FORBIDDEN,
        detail="Invalid WhatsApp webhook verification token",
    )


@app.post("/v1/whatsapp/ziplin/webhook")
@app.post("/v1/whatsapp/webhook")
async def whatsapp_webhook(request: Request):
    raw_body = await _read_bounded_whatsapp_body(request)
    if not verify_meta_signature(raw_body, request.headers.get("x-hub-signature-256")):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid WhatsApp webhook signature",
        )

    return await _enqueue_whatsapp_webhook(raw_body)


async def _read_bounded_whatsapp_body(request: Request) -> bytes:
    """Read a webhook body while keeping application memory strictly bounded."""
    declared_length = request.headers.get("content-length")
    if declared_length:
        try:
            parsed_length = int(declared_length)
        except ValueError as exc:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Invalid Content-Length header",
            ) from exc
        if parsed_length < 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Invalid Content-Length header",
            )
        if parsed_length > _WHATSAPP_WEBHOOK_MAX_BODY_BYTES:
            raise HTTPException(
                status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                detail="Webhook payload is too large",
            )

    body = bytearray()
    async for chunk in request.stream():
        if len(chunk) > _WHATSAPP_WEBHOOK_MAX_BODY_BYTES - len(body):
            raise HTTPException(
                status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                detail="Webhook payload is too large",
            )
        body.extend(chunk)
    return bytes(body)


def _whatsapp_routing_rejection(payload: dict, *, phone_number_id: str, waba_id: str) -> str | None:
    """Return a safe rejection reason unless every batched event belongs here."""
    if payload.get("object") != "whatsapp_business_account":
        return "invalid_webhook_routing"
    entries = payload.get("entry")
    if not isinstance(entries, list) or not entries:
        return "invalid_webhook_routing"

    expected_phone_id = phone_number_id.strip()
    expected_waba_id = waba_id.strip()
    if not expected_phone_id:
        return "different_phone_number"

    saw_event = False
    for entry in entries:
        if not isinstance(entry, dict):
            return "invalid_webhook_routing"
        if expected_waba_id:
            entry_id = entry.get("id")
            if not isinstance(entry_id, str) or entry_id.strip() != expected_waba_id:
                return "different_business_account"
        changes = entry.get("changes")
        if not isinstance(changes, list) or not changes:
            return "invalid_webhook_routing"
        for change in changes:
            if not isinstance(change, dict) or change.get("field") != "messages":
                return "invalid_webhook_routing"
            value = change.get("value")
            if not isinstance(value, dict):
                return "invalid_webhook_routing"
            messages = value.get("messages")
            statuses = value.get("statuses")
            if messages is not None and not isinstance(messages, list):
                return "invalid_webhook_routing"
            if statuses is not None and not isinstance(statuses, list):
                return "invalid_webhook_routing"
            events = [
                item
                for collection in (messages or [], statuses or [])
                for item in collection
            ]
            if not events or any(not isinstance(item, dict) for item in events):
                return "invalid_webhook_routing"
            saw_event = True
            metadata = value.get("metadata")
            if not isinstance(metadata, dict):
                return "different_phone_number"
            routed_phone_id = metadata.get("phone_number_id")
            if (
                not isinstance(routed_phone_id, str)
                or routed_phone_id.strip() != expected_phone_id
            ):
                return "different_phone_number"
    return None if saw_event else "invalid_webhook_routing"


def _status_only_whatsapp_payload(payload: dict) -> dict:
    """Return a minimal copy containing delivery/status receipts only.

    The dashboard kill switch pauses new student-message processing and all
    outbound sends. Meta status callbacks are safe to retain while paused and
    are needed to keep the read-only conversation timeline accurate. Removing
    messages and contacts before persistence also guarantees a mixed callback
    cannot be processed later after messaging is resumed.
    """
    filtered_entries: list[dict] = []
    for raw_entry in payload.get("entry") or []:
        if not isinstance(raw_entry, dict):
            continue
        filtered_changes: list[dict] = []
        for raw_change in raw_entry.get("changes") or []:
            if not isinstance(raw_change, dict):
                continue
            value = raw_change.get("value") or {}
            if not isinstance(value, dict):
                continue
            statuses = [item for item in value.get("statuses") or [] if isinstance(item, dict)]
            if not statuses:
                continue
            filtered_value: dict = {"statuses": statuses}
            metadata = value.get("metadata")
            if isinstance(metadata, dict):
                filtered_value["metadata"] = metadata
            filtered_change = {
                key: item for key, item in raw_change.items() if key != "value"
            }
            filtered_change["value"] = filtered_value
            filtered_changes.append(filtered_change)
        if filtered_changes:
            filtered_entry = {
                key: item for key, item in raw_entry.items() if key != "changes"
            }
            filtered_entry["changes"] = filtered_changes
            filtered_entries.append(filtered_entry)

    result = {key: item for key, item in payload.items() if key != "entry"}
    result["entry"] = filtered_entries
    return result


async def _enqueue_whatsapp_webhook(raw_body: bytes):
    if len(raw_body) > _WHATSAPP_WEBHOOK_MAX_BODY_BYTES:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Webhook payload is too large")
    try:
        payload = json.loads(raw_body)
    except (json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid JSON payload") from exc
    if not isinstance(payload, dict):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Webhook payload must be a JSON object")
    settings = get_settings()
    routing_rejection = _whatsapp_routing_rejection(
        payload,
        phone_number_id=settings.whatsapp_phone_number_id,
        waba_id=settings.whatsapp_business_account_id,
    )
    if routing_rejection:
        logger.warning(
            "WhatsApp webhook ignored before queueing reason=%s",
            routing_rejection,
        )
        return {
            "status": "ignored",
            "event_id": None,
            "reason": routing_rejection,
        }
    if not settings.whatsapp_messaging_enabled:
        # Keep Meta verification healthy while intentionally declining new
        # student messages. Delivery/read receipts remain durable so the admin
        # inbox does not become stale while replies are paused.
        payload = _status_only_whatsapp_payload(payload)
        if not payload["entry"]:
            return {"status": "paused", "event_id": None}
    event = await AdminStore().enqueue_whatsapp_webhook(payload)
    if event["status"] not in {"completed", "dead_letter"}:
        schedule_whatsapp_webhook(str(event["id"]))
    return {
        "status": "accepted" if settings.whatsapp_messaging_enabled else "paused",
        "event_id": event["id"],
    }


@app.post("/v1/whatsapp/relay")
async def whatsapp_webhook_relay(request: Request):
    """Legacy non-production compatibility path; Ziplin uses direct Meta delivery."""
    settings = get_settings()
    if settings.app_environment.casefold() == "production":
        raise HTTPException(
            status_code=status.HTTP_410_GONE,
            detail="Legacy WhatsApp relays are disabled in production",
        )
    configured = settings.whatsapp_relay_token.strip()
    if not configured:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="WhatsApp webhook relay is not configured",
        )
    supplied = (request.headers.get("x-northstar-relay-token") or "").strip()
    if not supplied or not secrets.compare_digest(supplied, configured):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid relay token")
    return await _enqueue_whatsapp_webhook(await _read_bounded_whatsapp_body(request))


@app.post("/v1/whatsapp/ziplin/relay")
async def ziplin_whatsapp_webhook_relay(request: Request):
    """Legacy non-production compatibility path; Ziplin uses direct Meta delivery."""
    settings = get_settings()
    if settings.app_environment.casefold() == "production":
        raise HTTPException(
            status_code=status.HTTP_410_GONE,
            detail="Legacy WhatsApp relays are disabled in production",
        )
    configured = settings.whatsapp_relay_token.strip()
    if not configured:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Ziplin WhatsApp webhook relay is not configured",
        )
    supplied = (request.headers.get("x-ziplin-relay-token") or "").strip()
    if not supplied or not secrets.compare_digest(supplied, configured):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid Ziplin relay token")
    return await _enqueue_whatsapp_webhook(await _read_bounded_whatsapp_body(request))


@app.post("/v1/admin/whatsapp/send-program-menu", dependencies=[Depends(require_admin_token)])
async def send_whatsapp_program_menu(request: WhatsAppMenuRequest):
    if not get_settings().whatsapp_messaging_enabled:
        raise HTTPException(status_code=409, detail="WhatsApp messaging is paused from the admin dashboard")
    client = WhatsAppClient()
    to = normalize_wa_id(request.to)
    courses = await get_enrolled_courses(Cache(), to)
    if not courses:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="WhatsApp user is not enrolled")
    ordered = ordered_courses(courses)
    if len(ordered) == 1:
        meta_response = await client.send_mode_menu(to, ordered[0])
        menu = "mode"
    else:
        meta_response = await client.send_program_menu(to, courses)
        menu = "course"
    return {"status": "sent", "to": to, "courses": ordered, "menu": menu, "meta_response": meta_response}


@app.post("/v1/admin/whatsapp/send-hi", dependencies=[Depends(require_admin_token)])
async def send_whatsapp_hi(request: WhatsAppStartRequest):
    settings = get_settings()
    if not settings.whatsapp_messaging_enabled:
        raise HTTPException(status_code=409, detail="WhatsApp messaging is paused from the admin dashboard")
    client = WhatsAppClient()
    to = normalize_wa_id(request.to)
    courses = await get_enrolled_courses(Cache(), to)
    if not courses:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="WhatsApp user is not enrolled")
    template_name = request.template_name or settings.whatsapp_start_template_name
    language_code = request.language_code or settings.whatsapp_start_template_language
    meta_response = await client.send_start_message(
        to=to,
        template_name=template_name,
        language_code=language_code,
    )
    return {
        "status": "sent",
        "to": to,
        "courses": ordered_courses(courses),
        "template_name": template_name,
        "language_code": language_code,
        "next_step": "Customer must reply to this WhatsApp template. Their reply opens their enrolled-course mode menu.",
        "meta_response": meta_response,
    }


def normalize_admin_enrollment_phone(phone: str) -> str:
    normalized_phone = normalize_wa_id(phone)
    if not 8 <= len(normalized_phone) <= 15:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Invalid international WhatsApp phone number",
        )
    return normalized_phone


@app.get(
    "/v1/admin/whatsapp/enrollments",
    dependencies=[Depends(require_admin_token)],
)
async def list_whatsapp_enrollments(
    search: str = Query(default="", max_length=30),
    limit: int = Query(default=50, ge=1, le=200),
    offset: int = Query(default=0, ge=0),
):
    try:
        items = await enrollment_roster()
    except EnrollmentWorkbookError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    normalized_search = normalize_wa_id(search) if search else ""
    if normalized_search:
        items = [item for item in items if normalized_search in item["phone"]]
    return {
        "items": items[offset : offset + limit],
        "total": len(items),
        "source": enrollment_source(),
        "limit": limit,
        "offset": offset,
    }


@app.post(
    "/v1/admin/whatsapp/enrollments/import",
    dependencies=[Depends(require_admin_token)],
)
async def import_whatsapp_enrollments(file: UploadFile = File(...)):
    filename = Path(file.filename or "enrollments.xlsx").name
    if Path(filename).suffix.casefold() != ".xlsx":
        raise HTTPException(status_code=415, detail="Upload an .xlsx Excel workbook")
    maximum = get_settings().admin_upload_max_bytes
    payload = await file.read(maximum + 1)
    await file.close()
    if not payload:
        raise HTTPException(status_code=422, detail="The enrollment workbook is empty")
    if len(payload) > maximum:
        raise HTTPException(status_code=413, detail="The enrollment workbook exceeds the upload limit")
    try:
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
        sheet = workbook["Enrollments"] if "Enrollments" in workbook.sheetnames else workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip().casefold() for value in next(values, ())]
        required = {"phone_number", "course", "active"}
        if not required.issubset(headers):
            raise ValueError("Required columns: phone_number, course, active")
        indexes = {name: headers.index(name) for name in headers if name}
        rows: list[tuple[str, str, bool, str, str]] = []
        for row_number, row in enumerate(values, start=2):
            phone = normalize_wa_id(str(row[indexes["phone_number"]] or ""))
            course = str(row[indexes["course"]] or "").strip().upper() or "CMA"
            active_text = str(row[indexes["active"]] or "").strip().casefold()
            if not phone and not course and not active_text:
                continue
            if not 8 <= len(phone) <= 15:
                raise ValueError(f"Row {row_number}: phone_number must contain 8 to 15 digits")
            if course not in COURSES:
                raise ValueError(f"Row {row_number}: unsupported course {course or '(blank)'}")
            if active_text not in {"", "yes", "no", "true", "false", "1", "0", "active", "inactive"}:
                raise ValueError(f"Row {row_number}: active must be blank, YES, or NO")
            active = active_text in {"", "yes", "true", "1", "active"}
            name = str(row[indexes["student_name"]] or "").strip() if "student_name" in indexes else ""
            notes = str(row[indexes["notes"]] or "").strip() if "notes" in indexes else ""
            rows.append((phone, course, active, name, notes))
            if len(rows) > 2000:
                raise ValueError("A bulk import can contain at most 2,000 enrollment rows")
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=422, detail="The Excel workbook could not be read") from exc
    finally:
        if "workbook" in locals():
            workbook.close()
    if not rows:
        raise HTTPException(status_code=422, detail="No enrollment rows were found")
    try:
        result = await bulk_set_enrollment_states(rows)
    except EnrollmentWorkbookError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    await AdminStore().record_audit(
        "enrollment.bulk_imported",
        "whatsapp_enrollment",
        filename,
        result,
    )
    return result


@app.put(
    "/v1/admin/whatsapp/enrollments/{phone}",
    response_model=WhatsAppEnrollmentResponse,
    dependencies=[Depends(require_admin_token)],
)
async def upsert_whatsapp_enrollment(phone: str, request: WhatsAppEnrollmentRequest):
    normalized_phone = normalize_admin_enrollment_phone(phone)
    cache = Cache()
    try:
        try:
            course = await set_enrolled_course(cache, normalized_phone, request.course)
        except EnrollmentSourceReadOnlyError as exc:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc)) from exc
        except EnrollmentWorkbookError as exc:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc)) from exc
        try:
            courses = ordered_courses(
                await get_enrolled_courses(
                    cache,
                    normalized_phone,
                    strict_external=True,
                )
            )
        except EnrollmentWorkbookError as exc:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc)) from exc
        if course not in courses:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail="The enrollment source did not persist the selected course.",
            )
        await AdminStore().record_audit(
            "enrollment.granted",
            "whatsapp_enrollment",
            normalized_phone,
            {"course": course, "source": enrollment_source()},
        )
        return WhatsAppEnrollmentResponse(
            phone=normalized_phone,
            course=course,
            courses=courses,
            source=enrollment_source(),
        )
    finally:
        await cache.aclose()


@app.get(
    "/v1/admin/whatsapp/enrollments/{phone}",
    response_model=WhatsAppEnrollmentResponse,
    dependencies=[Depends(require_admin_token)],
)
async def read_whatsapp_enrollment(phone: str):
    normalized_phone = normalize_admin_enrollment_phone(phone)
    cache = Cache()
    try:
        try:
            courses = ordered_courses(
                await get_enrolled_courses(
                    cache,
                    normalized_phone,
                    strict_external=True,
                )
            )
        except EnrollmentWorkbookError as exc:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc)) from exc
        return WhatsAppEnrollmentResponse(
            phone=normalized_phone,
            course=courses[0] if courses else None,
            courses=courses,
            source=enrollment_source(),
        )
    finally:
        await cache.aclose()


@app.delete(
    "/v1/admin/whatsapp/enrollments/{phone}/{course}",
    response_model=WhatsAppEnrollmentResponse,
    dependencies=[Depends(require_admin_token)],
)
async def delete_whatsapp_enrollment_course(phone: str, course: Course):
    normalized_phone = normalize_admin_enrollment_phone(phone)
    cache = Cache()
    try:
        try:
            remaining = ordered_courses(
                await remove_enrolled_course(cache, normalized_phone, course)
            )
        except EnrollmentSourceReadOnlyError as exc:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc)) from exc
        except EnrollmentWorkbookError as exc:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc)) from exc
        await AdminStore().record_audit(
            "enrollment.revoked",
            "whatsapp_enrollment",
            normalized_phone,
            {"course": course, "source": enrollment_source()},
        )
        return WhatsAppEnrollmentResponse(
            phone=normalized_phone,
            course=remaining[0] if remaining else None,
            courses=remaining,
            source=enrollment_source(),
        )
    finally:
        await cache.aclose()


@app.post("/v1/chat", response_model=ChatResponse)
async def chat(request: ChatRequest):
    service = MentorService()
    started = time.perf_counter()
    outcome = "success"
    try:
        return await service.answer(request)
    except Exception:
        outcome = "error"
        raise
    finally:
        await service.aclose()
        try:
            await AdminStore().record_usage(
                student_id=request.student_id,
                channel="web",
                course=request.course,
                mode=request.mode,
                status=outcome,
                latency_ms=round((time.perf_counter() - started) * 1000),
            )
        except Exception:
            logging.getLogger("uvicorn.error").exception("Could not record chat usage analytics")


@app.post("/v1/chat/stream")
async def chat_stream(request: ChatRequest):
    service = MentorService()

    async def stream_and_close():
        started = time.perf_counter()
        outcome = "success"
        try:
            async for chunk in service.stream_answer(request):
                yield chunk
        except Exception:
            outcome = "error"
            raise
        finally:
            await service.aclose()
            try:
                await AdminStore().record_usage(
                    student_id=request.student_id,
                    channel="web",
                    course=request.course,
                    mode=request.mode,
                    status=outcome,
                    latency_ms=round((time.perf_counter() - started) * 1000),
                )
            except Exception:
                logging.getLogger("uvicorn.error").exception("Could not record streamed chat usage analytics")

    return StreamingResponse(
        stream_and_close(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no",
            "Connection": "keep-alive",
        },
    )


@app.post("/v1/admin/ingest/files", response_model=IngestResult, dependencies=[Depends(require_admin_token)])
async def ingest_files(
    files: list[UploadFile] = File(...),
    course: Course = Form("CMA"),
    doc_type: DocumentType = Form("lesson"),
):
    created = await create_knowledge_documents(files=files, course=course, doc_type=doc_type)
    document_ids = [document["id"] for document in created["items"]]
    store = AdminStore()
    job = await store.create_training_job(document_ids)
    await run_training_job(job["id"])
    completed = await store.get_training_job(job["id"])
    if completed is None:
        raise HTTPException(status_code=500, detail="Knowledge indexing job record was lost")
    if completed["status"] not in {"completed", "partial"}:
        raise HTTPException(status_code=500, detail=completed.get("error_message") or "Knowledge indexing failed")

    return IngestResult(
        files=len(files),
        chunks=int(completed["total_chunks"]),
        collection=get_settings().qdrant_collection,
    )

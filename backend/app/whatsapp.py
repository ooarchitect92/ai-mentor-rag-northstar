import asyncio
import hashlib
import hmac
import logging
import os
import re
import threading
import time
from asyncio import to_thread
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any
from urllib.parse import quote

import httpx
from google.auth.transport.requests import Request as GoogleAuthRequest
from google.oauth2 import service_account
from openpyxl import Workbook, load_workbook

from .admin_store import AdminStore
from .cache import Cache
from .config import get_settings
from .gemini import GeminiService
from .mentor import MentorService
from .schemas import ChatRequest


logger = logging.getLogger("uvicorn.error")

COURSE_ORDER = ("CMA", "CPA", "CFA", "ACCA", "CS", "EA")
COURSES = set(COURSE_ORDER)
MODES = {"teach", "quiz", "revise", "job_hunt", "doubt_solving"}
LEVELS = {"beginner", "intermediate", "advanced"}
HI_TRIGGERS = {"hi", "hii", "hello", "hey", "start", "/start", "menu", "help", "program", "programs"}
FEEDBACK_CANCEL_TRIGGERS = {"cancel", "cancel feedback", "exit"}


def normalized_hi_trigger(text: str) -> str:
    """Normalize harmless trailing punctuation/emoji without matching sentences."""
    return re.sub(r"[\W_]+$", "", text.strip().casefold()).strip()


def is_hi_trigger(text: str) -> bool:
    return normalized_hi_trigger(text) in HI_TRIGGERS


class UnsupportedFeedbackMediaError(RuntimeError):
    """A permanent screenshot validation error that should not be retried."""


class WhatsAppInboundBusyError(RuntimeError):
    """A duplicate delivery is waiting for an earlier message attempt to expire."""

    def __init__(self, message_id: str, retry_after_seconds: int) -> None:
        super().__init__(f"WhatsApp message is still processing: {message_id}")
        self.retry_after_seconds = retry_after_seconds


class EnrollmentWorkbookError(RuntimeError):
    """The configured Excel enrollment source could not be updated safely."""


class EnrollmentSourceReadOnlyError(RuntimeError):
    """The configured external enrollment source is not writable by this app."""


PROGRAM_OPTIONS = [
    {
        "id": "program:cma",
        "label": "CMA",
        "course": "CMA",
        "description": "Certified Management Accountant",
    },
    {
        "id": "program:cpa",
        "label": "CPA",
        "course": "CPA",
        "description": "Certified Public Accountant",
    },
    {
        "id": "program:cfa",
        "label": "CFA",
        "course": "CFA",
        "description": "Chartered Financial Analyst",
    },
    {
        "id": "program:acca",
        "label": "ACCA",
        "course": "ACCA",
        "description": "Association of Chartered Certified Accountants",
    },
    {
        "id": "program:cs",
        "label": "CS",
        "course": "CS",
        "description": "Company Secretary",
    },
    {
        "id": "program:ea",
        "label": "EA",
        "course": "EA",
        "description": "Enrolled Agent",
    },
]

MODE_OPTIONS = [
    {
        "id": "mode:teach",
        "label": "Teach",
        "mode": "teach",
        "description": "Learn a concept step by step",
        "example": "Explain standard costing with a simple example.",
    },
    {
        "id": "mode:doubt_solving",
        "label": "Doubt Solving",
        "mode": "doubt_solving",
        "description": "Clear one specific doubt",
        "example": "Why is sales volume variance different from sales mix variance?",
    },
    {
        "id": "mode:quiz",
        "label": "Quiz",
        "mode": "quiz",
        "description": "Get practice questions",
        "example": "Give me 5 MCQs on variance analysis.",
    },
    {
        "id": "mode:revise",
        "label": "Revision",
        "mode": "revise",
        "description": "Revise quickly before exam",
        "example": "Revise marginal costing formulas.",
    },
    {
        "id": "mode:job_hunt",
        "label": "Job Hunt",
        "mode": "job_hunt",
        "description": "Resume, interview, and career help",
        "example": "Help me prepare for a finance interview using my course material.",
    },
]

MODE_BY_VALUE = {}
for option in MODE_OPTIONS:
    MODE_BY_VALUE[option["id"]] = option
    MODE_BY_VALUE[option["label"].lower()] = option
    MODE_BY_VALUE[option["mode"]] = option
MODE_BY_VALUE.update(
    {
        "1": MODE_OPTIONS[0],
        "2": MODE_OPTIONS[1],
        "doubt": MODE_OPTIONS[1],
        "doubt solving": MODE_OPTIONS[1],
        "doubt-solving": MODE_OPTIONS[1],
        "3": MODE_OPTIONS[2],
        "4": MODE_OPTIONS[3],
        "revision": MODE_OPTIONS[3],
        "revise": MODE_OPTIONS[3],
        "5": MODE_OPTIONS[4],
        "job": MODE_OPTIONS[4],
        "job hunt": MODE_OPTIONS[4],
        "job-hunt": MODE_OPTIONS[4],
    }
)

PROGRAM_BY_VALUE = {}
for option in PROGRAM_OPTIONS:
    PROGRAM_BY_VALUE[option["id"]] = option
    PROGRAM_BY_VALUE[option["label"].lower()] = option
    PROGRAM_BY_VALUE[option["course"].lower()] = option
PROGRAM_BY_VALUE.update({str(index): option for index, option in enumerate(PROGRAM_OPTIONS, start=1)})
PROGRAM_BY_VALUE.update(
    {
        "cma us": PROGRAM_OPTIONS[0],
        "cma usa": PROGRAM_OPTIONS[0],
        "cpa us": PROGRAM_OPTIONS[1],
        "cpa usa": PROGRAM_OPTIONS[1],
        "company secretary": next(option for option in PROGRAM_OPTIONS if option["course"] == "CS"),
    }
)


def normalize_wa_id(value: str, default_country_code: str = "91") -> str:
    digits = "".join(char for char in value if char.isdigit())
    if len(digits) == 10:
        return f"{default_country_code}{digits}"
    return digits


def enrollment_key(sender: str) -> str:
    return f"whatsapp:enrollment:{normalize_wa_id(sender)}"


def _parse_course_list(raw: str) -> set[str]:
    normalized = raw.replace("+", "|").replace(",", "|")
    return {value.strip().upper() for value in normalized.split("|") if value.strip().upper() in COURSES}


def configured_enrollments(raw: str) -> dict[str, set[str]]:
    enrollments: dict[str, set[str]] = {}
    for item in raw.replace(";", ",").split(","):
        phone, separator, course = item.strip().partition(":")
        normalized_phone = normalize_wa_id(phone)
        normalized_courses = _parse_course_list(course)
        if separator and normalized_phone and normalized_courses:
            enrollments.setdefault(normalized_phone, set()).update(normalized_courses)
    return enrollments


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_enrollment_rows(rows: list[list[Any]]) -> dict[str, set[str]]:
    if not rows:
        return {}
    headers = [_cell_text(value).lower() for value in rows[0]]
    required = {"phone_number", "course", "active"}
    if not required.issubset(headers):
        raise ValueError("Enrollment sheet is missing required columns")
    indexes = {name: headers.index(name) for name in required}

    states: dict[tuple[str, str], bool] = {}
    phones: set[str] = set()
    for row in rows[1:]:
        padded = list(row) + [""] * max(0, len(headers) - len(row))
        phone = normalize_wa_id(_cell_text(padded[indexes["phone_number"]]))
        if len(phone) < 8:
            continue
        phones.add(phone)
        # Enrollment is intentionally CMA-first: adding a valid phone number
        # to the maintained sheet is enough to grant CMA access. Operators can
        # still specify another supported course or explicitly enter NO to
        # deny/revoke the row.
        course = _cell_text(padded[indexes["course"]]).upper() or "CMA"
        active = _cell_text(padded[indexes["active"]]).upper()
        is_active = active in {"", "YES", "TRUE", "1", "ACTIVE"}
        if course in COURSES:
            states[(phone, course)] = is_active

    enrollments: dict[str, set[str]] = {phone: set() for phone in phones}
    for (phone, course), is_active in states.items():
        if is_active:
            enrollments[phone].add(course)
    return enrollments


@lru_cache(maxsize=8)
def _load_enrollment_workbook(path_value: str, modified_ns: int) -> dict[str, set[str]]:
    del modified_ns  # Included in the cache key so saving the workbook reloads it.
    workbook = load_workbook(path_value, read_only=True, data_only=True)
    try:
        if "Enrollments" not in workbook.sheetnames:
            raise ValueError("Enrollment workbook must contain an Enrollments sheet")
        sheet = workbook["Enrollments"]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip().lower() for value in next(rows, ())]
        required = {"phone_number", "course", "active"}
        if not required.issubset(headers):
            raise ValueError("Enrollment workbook is missing required columns")

        return _parse_enrollment_rows([headers, *[list(row) for row in rows]])
    finally:
        workbook.close()


def workbook_enrollments(path_value: str) -> dict[str, set[str]]:
    if not path_value:
        return {}
    path = Path(path_value)
    try:
        return _load_enrollment_workbook(str(path), path.stat().st_mtime_ns)
    except Exception:
        logger.exception("Could not load WhatsApp enrollment workbook %s", path)
        return {}


def _strict_workbook_enrollments(path_value: str) -> dict[str, set[str]]:
    path = Path(path_value)
    try:
        return _load_enrollment_workbook(str(path), path.stat().st_mtime_ns)
    except FileNotFoundError as exc:
        raise EnrollmentWorkbookError("The configured Excel enrollment workbook was not found") from exc
    except EnrollmentWorkbookError:
        raise
    except Exception as exc:
        raise EnrollmentWorkbookError("The configured Excel enrollment workbook could not be read") from exc


_enrollment_workbook_write_lock = threading.RLock()
_ENROLLMENT_WORKBOOK_HEADERS = (
    "phone_number",
    "course",
    "active",
    "student_name",
    "notes",
)


def enrollment_source() -> str:
    settings = get_settings()
    if settings.whatsapp_enrollments_google_sheet_id.strip():
        return "google_sheet"
    if settings.whatsapp_enrollments_file.strip():
        return "excel"
    return "redis"


def _ensure_enrollment_sheet(workbook: Workbook):
    if "Enrollments" in workbook.sheetnames:
        sheet = workbook["Enrollments"]
    elif len(workbook.sheetnames) == 1 and workbook.active.max_row == 1 and workbook.active["A1"].value is None:
        sheet = workbook.active
        sheet.title = "Enrollments"
    else:
        raise EnrollmentWorkbookError("Excel workbook must contain an Enrollments sheet")

    headers = [
        _cell_text(sheet.cell(row=1, column=column).value).lower()
        for column in range(1, max(sheet.max_column, len(_ENROLLMENT_WORKBOOK_HEADERS)) + 1)
    ]
    if not any(headers):
        for column, header in enumerate(_ENROLLMENT_WORKBOOK_HEADERS, start=1):
            sheet.cell(row=1, column=column, value=header)
        headers = list(_ENROLLMENT_WORKBOOK_HEADERS)
        sheet.freeze_panes = "A2"

    required = {"phone_number", "course", "active"}
    if not required.issubset(headers):
        raise EnrollmentWorkbookError(
            "Enrollments sheet must contain phone_number, course, and active columns"
        )
    return sheet, {name: headers.index(name) + 1 for name in required}


def _set_workbook_enrollment_sync(
    path_value: str,
    sender: str,
    course: str,
    active: bool,
    seed_courses: set[str] | None = None,
) -> set[str]:
    path = Path(path_value)
    if path.suffix.casefold() != ".xlsx":
        raise EnrollmentWorkbookError("WHATSAPP_ENROLLMENTS_FILE must point to an .xlsx workbook")
    normalized_sender = normalize_wa_id(sender)
    normalized_course = course.upper()
    if not 8 <= len(normalized_sender) <= 15 or normalized_course not in COURSES:
        raise EnrollmentWorkbookError("Invalid phone number or course for the enrollment workbook")

    temporary_path: Path | None = None
    workbook = None
    with _enrollment_workbook_write_lock:
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            original_signature = (
                (path.stat().st_mtime_ns, path.stat().st_size)
                if path.exists()
                else None
            )
            workbook = load_workbook(path) if path.exists() else Workbook()
            sheet, indexes = _ensure_enrollment_sheet(workbook)
            matching_rows: list[int] = []
            phone_rows: list[int] = []
            last_data_row = 1
            for row_number in range(2, sheet.max_row + 1):
                row_phone = normalize_wa_id(
                    _cell_text(sheet.cell(row=row_number, column=indexes["phone_number"]).value)
                )
                row_course = _cell_text(
                    sheet.cell(row=row_number, column=indexes["course"]).value
                ).upper()
                row_active = _cell_text(
                    sheet.cell(row=row_number, column=indexes["active"]).value
                )
                if row_phone or row_course or row_active:
                    last_data_row = row_number
                if row_phone == normalized_sender:
                    phone_rows.append(row_number)
                    if row_course == normalized_course:
                        matching_rows.append(row_number)

            changed = False
            desired_active = "YES" if active else "NO"
            if not phone_rows and seed_courses is not None:
                seeded_states = {
                    course_name: True
                    for course_name in seed_courses
                    if course_name in COURSES
                }
                seeded_states[normalized_course] = active
                for course_name in COURSE_ORDER:
                    if course_name not in seeded_states:
                        continue
                    last_data_row += 1
                    row_number = last_data_row
                    sheet.cell(row=row_number, column=indexes["phone_number"], value=normalized_sender)
                    sheet.cell(row=row_number, column=indexes["phone_number"]).number_format = "@"
                    sheet.cell(row=row_number, column=indexes["course"], value=course_name)
                    sheet.cell(
                        row=row_number,
                        column=indexes["active"],
                        value="YES" if seeded_states[course_name] else "NO",
                    )
                changed = True
            elif matching_rows:
                row_number = matching_rows[-1]
                active_cell = sheet.cell(row=row_number, column=indexes["active"])
                if _cell_text(active_cell.value).upper() != desired_active:
                    active_cell.value = desired_active
                    changed = True
            elif active:
                last_data_row += 1
                row_number = last_data_row
                sheet.cell(row=row_number, column=indexes["phone_number"], value=normalized_sender)
                sheet.cell(row=row_number, column=indexes["phone_number"]).number_format = "@"
                sheet.cell(row=row_number, column=indexes["course"], value=normalized_course)
                sheet.cell(row=row_number, column=indexes["active"], value="YES")
                changed = True

            if changed:
                last_column = max(sheet.max_column, len(_ENROLLMENT_WORKBOOK_HEADERS))
                sheet.auto_filter.ref = (
                    f"A1:{sheet.cell(row=1, column=last_column).column_letter}{last_data_row}"
                )
                for table in sheet.tables.values():
                    if str(table.ref).upper().startswith("A1:"):
                        table.ref = (
                            f"A1:{sheet.cell(row=1, column=last_column).column_letter}{last_data_row}"
                        )

                with NamedTemporaryFile(
                    dir=path.parent,
                    prefix=f".{path.stem}.",
                    suffix=".xlsx",
                    delete=False,
                ) as temporary:
                    temporary_path = Path(temporary.name)
                workbook.save(temporary_path)
                workbook.close()
                workbook = None
                with temporary_path.open("r+b") as saved_workbook:
                    saved_workbook.flush()
                    os.fsync(saved_workbook.fileno())

                validation = load_workbook(temporary_path, read_only=True, data_only=True)
                try:
                    _ensure_enrollment_sheet(validation)
                finally:
                    validation.close()
                current_signature = (
                    (path.stat().st_mtime_ns, path.stat().st_size)
                    if path.exists()
                    else None
                )
                if current_signature != original_signature:
                    raise EnrollmentWorkbookError(
                        "The Excel enrollment workbook changed while saving. Refresh and try again."
                    )
                os.replace(temporary_path, path)
                temporary_path = None
                _load_enrollment_workbook.cache_clear()
        except EnrollmentWorkbookError:
            raise
        except PermissionError as exc:
            raise EnrollmentWorkbookError(
                "The Excel enrollment workbook is open or not writable. Close it and try again."
            ) from exc
        except OSError as exc:
            raise EnrollmentWorkbookError(
                "The Excel enrollment workbook could not be saved. Check its path and permissions."
            ) from exc
        except Exception as exc:
            raise EnrollmentWorkbookError(
                "The Excel enrollment workbook could not be updated safely."
            ) from exc
        finally:
            if workbook is not None:
                workbook.close()
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)

        _load_enrollment_workbook.cache_clear()
        return set(_strict_workbook_enrollments(str(path)).get(normalized_sender, set()))


async def set_workbook_enrollment(
    path_value: str,
    sender: str,
    course: str,
    *,
    active: bool,
    seed_courses: set[str] | None = None,
) -> set[str]:
    return await to_thread(
        _set_workbook_enrollment_sync,
        path_value,
        sender,
        course,
        active,
        seed_courses,
    )


_google_sheet_cache: dict[str, set[str]] | None = None
_google_sheet_cache_key = ""
_google_sheet_cache_expires_at = 0.0


def _fetch_google_sheet_enrollments(
    spreadsheet_id: str,
    sheet_range: str,
    service_account_file: str,
) -> dict[str, set[str]]:
    credentials = service_account.Credentials.from_service_account_file(
        service_account_file,
        scopes=["https://www.googleapis.com/auth/spreadsheets"],
    )
    credentials.refresh(GoogleAuthRequest())
    encoded_range = quote(sheet_range, safe="")
    url = (
        f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}"
        f"/values/{encoded_range}"
    )
    response = httpx.get(
        url,
        headers={"Authorization": f"Bearer {credentials.token}"},
        timeout=15.0,
    )
    response.raise_for_status()
    payload = response.json()
    return _parse_enrollment_rows(payload.get("values", []))


def _append_google_sheet_enrollment_sync(
    spreadsheet_id: str,
    sheet_range: str,
    service_account_file: str,
    sender: str,
    course: str,
    active: bool,
) -> None:
    credentials = service_account.Credentials.from_service_account_file(
        service_account_file,
        scopes=["https://www.googleapis.com/auth/spreadsheets"],
    )
    credentials.refresh(GoogleAuthRequest())
    encoded_range = quote(sheet_range, safe="")
    url = (
        f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}"
        f"/values/{encoded_range}:append"
    )
    response = httpx.post(
        url,
        params={"valueInputOption": "RAW", "insertDataOption": "INSERT_ROWS"},
        headers={"Authorization": f"Bearer {credentials.token}"},
        json={
            "majorDimension": "ROWS",
            "values": [[sender, course, "YES" if active else "NO", "", "Updated from admin dashboard"]],
        },
        timeout=20.0,
    )
    response.raise_for_status()


def _append_google_sheet_rows_sync(
    spreadsheet_id: str,
    sheet_range: str,
    service_account_file: str,
    rows: list[tuple[str, str, bool, str, str]],
) -> None:
    credentials = service_account.Credentials.from_service_account_file(
        service_account_file,
        scopes=["https://www.googleapis.com/auth/spreadsheets"],
    )
    credentials.refresh(GoogleAuthRequest())
    encoded_range = quote(sheet_range, safe="")
    response = httpx.post(
        f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}/values/{encoded_range}:append",
        params={"valueInputOption": "RAW", "insertDataOption": "INSERT_ROWS"},
        headers={"Authorization": f"Bearer {credentials.token}"},
        json={
            "majorDimension": "ROWS",
            "values": [
                [phone, course, "YES" if active else "NO", name, notes]
                for phone, course, active, name, notes in rows
            ],
        },
        timeout=30.0,
    )
    response.raise_for_status()


async def bulk_set_enrollment_states(
    rows: list[tuple[str, str, bool, str, str]],
) -> dict[str, int | str]:
    global _google_sheet_cache, _google_sheet_cache_key, _google_sheet_cache_expires_at
    settings = get_settings()
    source = enrollment_source()
    normalized_rows = [
        (normalize_wa_id(phone), course.upper(), active, name[:200], notes[:500])
        for phone, course, active, name, notes in rows
    ]
    if source == "google_sheet":
        try:
            await to_thread(
                _append_google_sheet_rows_sync,
                settings.whatsapp_enrollments_google_sheet_id.strip(),
                settings.whatsapp_enrollments_google_sheet_range,
                settings.google_service_account_file.strip(),
                normalized_rows,
            )
        except Exception as exc:
            logger.exception("Could not bulk write WhatsApp enrollments to Google Sheets")
            raise EnrollmentWorkbookError(
                "Google Sheet bulk update failed. Confirm Editor sharing and try again."
            ) from exc
        _google_sheet_cache = None
        _google_sheet_cache_key = ""
        _google_sheet_cache_expires_at = 0.0
        await google_sheet_enrollments()
    else:
        cache = Cache()
        try:
            for phone, course, active, _name, _notes in normalized_rows:
                if active:
                    await set_enrolled_course(cache, phone, course)
                else:
                    await remove_enrolled_course(cache, phone, course)
        finally:
            await cache.aclose()
    return {
        "source": source,
        "rows": len(normalized_rows),
        "granted": sum(1 for row in normalized_rows if row[2]),
        "revoked": sum(1 for row in normalized_rows if not row[2]),
    }


async def set_google_sheet_enrollment(sender: str, course: str, *, active: bool) -> set[str]:
    global _google_sheet_cache, _google_sheet_cache_key, _google_sheet_cache_expires_at
    settings = get_settings()
    spreadsheet_id = settings.whatsapp_enrollments_google_sheet_id.strip()
    credentials_file = settings.google_service_account_file.strip()
    if not spreadsheet_id or not credentials_file:
        raise EnrollmentWorkbookError("Google Sheets enrollment credentials are incomplete")
    try:
        await to_thread(
            _append_google_sheet_enrollment_sync,
            spreadsheet_id,
            settings.whatsapp_enrollments_google_sheet_range,
            credentials_file,
            normalize_wa_id(sender),
            course.upper(),
            active,
        )
    except Exception as exc:
        logger.exception("Could not write WhatsApp enrollment to Google Sheets")
        raise EnrollmentWorkbookError(
            "Google Sheet could not be updated. Share it with the service account as Editor and retry."
        ) from exc
    _google_sheet_cache = None
    _google_sheet_cache_key = ""
    _google_sheet_cache_expires_at = 0.0
    refreshed = await google_sheet_enrollments()
    return set((refreshed or {}).get(normalize_wa_id(sender), set()))


async def enrollment_roster() -> list[dict[str, Any]]:
    settings = get_settings()
    source = enrollment_source()
    if source == "google_sheet":
        values = await google_sheet_enrollments() or {}
    elif source == "excel":
        values = await to_thread(_strict_workbook_enrollments, settings.whatsapp_enrollments_file)
    else:
        values = configured_enrollments(settings.whatsapp_enrollments)
    return [
        {"phone": phone, "courses": ordered_courses(set(courses)), "source": source}
        for phone, courses in sorted(values.items())
        if courses
    ]


async def google_sheet_enrollments() -> dict[str, set[str]] | None:
    """Return None when disabled; otherwise the Sheet is the sole access authority."""
    global _google_sheet_cache, _google_sheet_cache_key, _google_sheet_cache_expires_at

    settings = get_settings()
    spreadsheet_id = settings.whatsapp_enrollments_google_sheet_id.strip()
    credentials_file = settings.google_service_account_file.strip()
    if not spreadsheet_id:
        return None
    if not credentials_file:
        logger.error("Google Sheet enrollment is enabled without GOOGLE_SERVICE_ACCOUNT_FILE")
        return {}

    cache_key = f"{spreadsheet_id}:{settings.whatsapp_enrollments_google_sheet_range}:{credentials_file}"
    now = time.monotonic()
    if _google_sheet_cache_key == cache_key and now < _google_sheet_cache_expires_at:
        return _google_sheet_cache or {}

    try:
        enrollments = await to_thread(
            _fetch_google_sheet_enrollments,
            spreadsheet_id,
            settings.whatsapp_enrollments_google_sheet_range,
            credentials_file,
        )
    except Exception:
        logger.exception("Could not refresh WhatsApp enrollments from Google Sheets")
        if _google_sheet_cache_key == cache_key and _google_sheet_cache is not None:
            return _google_sheet_cache
        return {}

    _google_sheet_cache = enrollments
    _google_sheet_cache_key = cache_key
    _google_sheet_cache_expires_at = now + max(
        10,
        settings.whatsapp_enrollments_google_refresh_seconds,
    )
    return enrollments


async def get_enrolled_courses(
    cache: Cache,
    sender: str,
    *,
    strict_external: bool = False,
) -> set[str]:
    normalized_sender = normalize_wa_id(sender)
    settings = get_settings()
    google_spreadsheet = await google_sheet_enrollments()
    if google_spreadsheet is not None:
        return set(google_spreadsheet.get(normalized_sender, set()))
    workbook_path = settings.whatsapp_enrollments_file.strip()
    if workbook_path:
        try:
            spreadsheet = await to_thread(_strict_workbook_enrollments, workbook_path)
        except EnrollmentWorkbookError:
            if strict_external:
                raise
            logger.exception("Configured Excel enrollment source is unavailable")
            return set()
        # A configured workbook is the sole local authority, including missing
        # and explicitly disabled students. Redis cannot silently re-grant them.
        if normalized_sender in spreadsheet:
            return set(spreadsheet[normalized_sender])

        # One-time compatibility migration for enrollments saved by older
        # dashboard versions, which wrote only to Redis when a workbook row was missing.
        legacy_value = await cache.get_text(enrollment_key(normalized_sender))
        if legacy_value is not None:
            legacy_courses = _parse_course_list(legacy_value)
            if not legacy_courses:
                await cache.delete(enrollment_key(normalized_sender))
                return set()
            first_course = next(
                course_name for course_name in COURSE_ORDER if course_name in legacy_courses
            )
            try:
                migrated = await set_workbook_enrollment(
                    workbook_path,
                    normalized_sender,
                    first_course,
                    active=True,
                    seed_courses=legacy_courses,
                )
            except EnrollmentWorkbookError:
                if strict_external:
                    raise
                logger.exception(
                    "Could not migrate legacy Redis enrollment to Excel sender=%s",
                    normalized_sender,
                )
                return set()
            await cache.delete(enrollment_key(normalized_sender))
            return migrated
        return set()
    stored = await cache.get_text(enrollment_key(normalized_sender))
    if stored is not None:
        return _parse_course_list(stored)
    return set(configured_enrollments(settings.whatsapp_enrollments).get(normalized_sender, set()))


async def get_enrolled_course(cache: Cache, sender: str) -> str | None:
    """Backward-compatible primary course lookup; multi-course flows use get_enrolled_courses."""
    courses = await get_enrolled_courses(cache, sender)
    selected = await cache.get_text(program_session_key(normalize_wa_id(sender)))
    if selected in courses:
        return selected
    return next((course for course in COURSE_ORDER if course in courses), None)


async def set_enrolled_course(cache: Cache, sender: str, course: str) -> str:
    normalized_sender = normalize_wa_id(sender)
    normalized_course = course.upper()
    if normalized_course not in COURSES:
        raise ValueError("Unsupported enrollment course")
    settings = get_settings()
    google_enabled = bool(settings.whatsapp_enrollments_google_sheet_id.strip())
    workbook_path = settings.whatsapp_enrollments_file.strip()
    if google_enabled:
        existing = await set_google_sheet_enrollment(
            normalized_sender,
            normalized_course,
            active=True,
        )
    elif workbook_path:
        workbook_snapshot = (
            await to_thread(_strict_workbook_enrollments, workbook_path)
            if Path(workbook_path).exists()
            else {}
        )
        seed_courses: set[str] | None = None
        if normalized_sender not in workbook_snapshot:
            legacy_value = await cache.get_text(enrollment_key(normalized_sender))
            if legacy_value is not None:
                seed_courses = _parse_course_list(legacy_value)
            else:
                seed_courses = set(
                    configured_enrollments(settings.whatsapp_enrollments).get(
                        normalized_sender,
                        set(),
                    )
                )
        existing = await set_workbook_enrollment(
            workbook_path,
            normalized_sender,
            normalized_course,
            active=True,
            seed_courses=seed_courses,
        )
        await cache.delete(enrollment_key(normalized_sender))
    else:
        existing = await get_enrolled_courses(cache, normalized_sender)
        existing.add(normalized_course)
    ordered = [course_name for course_name in COURSE_ORDER if course_name in existing]
    if not workbook_path and not google_enabled:
        await cache.set_text(
            enrollment_key(normalized_sender),
            ",".join(ordered),
            ttl_seconds=315360000,
        )
    await cache.set_text(
        program_session_key(normalized_sender),
        normalized_course,
        ttl_seconds=get_settings().whatsapp_session_ttl_seconds,
    )
    await cache.delete(mode_session_key(normalized_sender))
    return normalized_course


async def remove_enrolled_course(cache: Cache, sender: str, course: str) -> set[str]:
    normalized_sender = normalize_wa_id(sender)
    normalized_course = course.upper()
    if normalized_course not in COURSES:
        raise ValueError("Unsupported enrollment course")
    settings = get_settings()
    google_enabled = bool(settings.whatsapp_enrollments_google_sheet_id.strip())
    workbook_path = settings.whatsapp_enrollments_file.strip()
    if google_enabled:
        existing = await set_google_sheet_enrollment(
            normalized_sender,
            normalized_course,
            active=False,
        )
    elif workbook_path:
        workbook_snapshot = (
            await to_thread(_strict_workbook_enrollments, workbook_path)
            if Path(workbook_path).exists()
            else {}
        )
        seed_courses: set[str] | None = None
        if normalized_sender not in workbook_snapshot:
            legacy_value = await cache.get_text(enrollment_key(normalized_sender))
            if legacy_value is not None:
                seed_courses = _parse_course_list(legacy_value)
            else:
                seed_courses = set(
                    configured_enrollments(settings.whatsapp_enrollments).get(
                        normalized_sender,
                        set(),
                    )
                )
        existing = await set_workbook_enrollment(
            workbook_path,
            normalized_sender,
            normalized_course,
            active=False,
            seed_courses=seed_courses,
        )
        await cache.delete(enrollment_key(normalized_sender))
    else:
        existing = await get_enrolled_courses(cache, normalized_sender)
        existing.discard(normalized_course)
    ordered = [course_name for course_name in COURSE_ORDER if course_name in existing]
    if not workbook_path and not google_enabled:
        await cache.set_text(
            enrollment_key(normalized_sender),
            ",".join(ordered) if ordered else "NONE",
            ttl_seconds=315360000,
        )
    selected = await cache.get_text(program_session_key(normalized_sender))
    if selected == normalized_course:
        await cache.delete(program_session_key(normalized_sender))
        await cache.delete(mode_session_key(normalized_sender))
    return existing


@dataclass(frozen=True)
class WhatsAppIncomingMessage:
    message_id: str
    sender: str
    text: str
    message_type: str
    phone_number_id: str | None = None
    profile_name: str | None = None
    media_id: str | None = None
    media_mime_type: str | None = None


@dataclass(frozen=True)
class WhatsAppStatusEvent:
    message_id: str
    status: str
    recipient_id: str
    timestamp: str
    conversation_id: str | None
    errors: list[dict[str, Any]]


def verify_meta_signature(raw_body: bytes, signature_header: str | None) -> bool:
    settings = get_settings()
    app_secret = settings.whatsapp_app_secret or settings.meta_app_secret
    if not app_secret:
        return settings.whatsapp_use_mock or settings.app_environment == "test"

    if not signature_header or not signature_header.startswith("sha256="):
        return False

    expected = "sha256=" + hmac.new(
        app_secret.encode("utf-8"),
        raw_body,
        hashlib.sha256,
    ).hexdigest()
    return hmac.compare_digest(expected, signature_header)


def extract_incoming_messages(payload: dict[str, Any]) -> list[WhatsAppIncomingMessage]:
    messages: list[WhatsAppIncomingMessage] = []

    for entry in payload.get("entry", []):
        for change in entry.get("changes", []):
            value = change.get("value", {})
            phone_number_id = str((value.get("metadata") or {}).get("phone_number_id") or "").strip() or None
            contacts_by_id = {
                contact.get("wa_id"): (contact.get("profile") or {}).get("name")
                for contact in value.get("contacts", [])
            }

            for raw_message in value.get("messages", []):
                sender = str(raw_message.get("from", ""))
                message_id = str(raw_message.get("id", ""))
                message_type = str(raw_message.get("type", ""))
                text = ""
                media_id = None
                media_mime_type = None
                if message_type == "text":
                    text = str((raw_message.get("text") or {}).get("body", "")).strip()
                elif message_type == "image":
                    image = raw_message.get("image") or {}
                    text = str(image.get("caption") or "").strip()
                    media_id = str(image.get("id") or "") or None
                    media_mime_type = str(image.get("mime_type") or "") or None
                elif message_type == "interactive":
                    interactive = raw_message.get("interactive") or {}
                    if interactive.get("type") == "list_reply":
                        reply = interactive.get("list_reply") or {}
                        text = str(reply.get("id") or reply.get("title") or "").strip()
                    elif interactive.get("type") == "button_reply":
                        reply = interactive.get("button_reply") or {}
                        text = str(reply.get("id") or reply.get("title") or "").strip()
                elif message_type == "button":
                    button = raw_message.get("button") or {}
                    text = str(button.get("payload") or button.get("text") or "").strip()

                if sender and message_id:
                    messages.append(
                        WhatsAppIncomingMessage(
                            message_id=message_id,
                            sender=sender,
                            text=text,
                            message_type=message_type,
                            phone_number_id=phone_number_id,
                            profile_name=contacts_by_id.get(sender),
                            media_id=media_id,
                            media_mime_type=media_mime_type,
                        )
                    )

    return messages


def extract_status_events(payload: dict[str, Any]) -> list[WhatsAppStatusEvent]:
    events: list[WhatsAppStatusEvent] = []

    for entry in payload.get("entry", []):
        for change in entry.get("changes", []):
            value = change.get("value", {})
            for raw_status in value.get("statuses", []):
                message_id = str(raw_status.get("id", ""))
                status = str(raw_status.get("status", ""))
                recipient_id = str(raw_status.get("recipient_id", ""))
                timestamp = str(raw_status.get("timestamp", ""))
                conversation = raw_status.get("conversation") or {}
                errors = raw_status.get("errors") or []

                if message_id and status:
                    events.append(
                        WhatsAppStatusEvent(
                            message_id=message_id,
                            status=status,
                            recipient_id=recipient_id,
                            timestamp=timestamp,
                            conversation_id=conversation.get("id"),
                            errors=errors if isinstance(errors, list) else [],
                        )
                    )

    return events


def split_whatsapp_text(text: str, max_chars: int) -> list[str]:
    text = text.strip()
    if not text:
        return []

    chunks: list[str] = []
    remaining = text
    while len(remaining) > max_chars:
        split_at = max(
            remaining.rfind("\n\n", 0, max_chars),
            remaining.rfind("\n", 0, max_chars),
            remaining.rfind(" ", 0, max_chars),
        )
        if split_at < max_chars * 0.5:
            split_at = max_chars

        chunks.append(remaining[:split_at].strip())
        remaining = remaining[split_at:].strip()

    if remaining:
        chunks.append(remaining)
    return chunks


def mode_session_key(sender: str) -> str:
    return f"whatsapp:mode:{sender}"


def program_session_key(sender: str) -> str:
    return f"whatsapp:program:{sender}"


def conversation_started_key(sender: str) -> str:
    return f"whatsapp:started:{sender}"


def feedback_session_key(sender: str) -> str:
    return f"whatsapp:feedback:{normalize_wa_id(sender)}"


def split_feedback_marker(text: str) -> tuple[bool, str]:
    """Recognize a deliberate FEEDBACK marker without stealing normal questions."""
    stripped = text.strip()
    if stripped.casefold() == "feedback":
        return True, ""
    bracketed = re.match(r"^\[feedback\]\s*(?::|-)?\s*(.*)$", stripped, flags=re.IGNORECASE | re.DOTALL)
    if bracketed is not None:
        return True, bracketed.group(1).strip()
    delimited = re.match(r"^feedback\s*(?::|-|\r?\n)\s*(.*)$", stripped, flags=re.IGNORECASE | re.DOTALL)
    if delimited is not None:
        return True, delimited.group(1).strip()
    return False, ""


def text_menu() -> str:
    return program_menu()


def ordered_courses(courses: set[str] | None = None) -> list[str]:
    allowed = courses if courses is not None else COURSES
    return [course for course in COURSE_ORDER if course in allowed]


def program_menu(courses: set[str] | None = None) -> str:
    available = ordered_courses(courses)
    course_lines = "\n".join(f"- {course}" for course in available)
    return (
        "*Choose your enrolled course*\n\n"
        f"{course_lines}\n\n"
        "Reply with the course name or tap the course option."
    )


def mode_menu(course: str) -> str:
    return (
        f"*{course} selected*\n"
        "Now choose how I should help you:\n\n"
        "1. Teach\n"
        "2. Doubt Solving\n"
        "3. Quiz\n"
        "4. Revision\n"
        "5. Job Hunt\n\n"
        "Reply with 1-5 or tap the menu option.\n"
        "Type *feedback* anytime to report a change or error."
    )


def normalize_program_selection(text: str) -> tuple[dict[str, str] | None, str]:
    cleaned = " ".join(text.strip().split())
    if not cleaned:
        return None, ""

    lower = cleaned.lower()
    for option in PROGRAM_OPTIONS:
        for prefix in (f"{option['label'].lower()}:", f"{option['course'].lower()}:", f"{option['id']}:"):
            if lower.startswith(prefix):
                return option, cleaned[len(prefix) :].strip()

    return PROGRAM_BY_VALUE.get(lower), ""


def normalize_mode_selection(text: str) -> tuple[dict[str, str] | None, str]:
    cleaned = " ".join(text.strip().split())
    if not cleaned:
        return None, ""

    lower = cleaned.lower()
    for option in MODE_OPTIONS:
        for prefix in (f"{option['label'].lower()}:", f"{option['mode']}:", f"{option['id']}:"):
            if lower.startswith(prefix):
                return option, cleaned[len(prefix) :].strip()

    return MODE_BY_VALUE.get(lower), ""


def mode_prompt(option: dict[str, str]) -> str:
    return (
        f"*{option['label']} selected*\n"
        "Send your question now.\n\n"
        f"Example: {option['example']}\n\n"
        "Type *menu* anytime to change learning mode."
    )


def format_whatsapp_answer(answer: str, course: str, option: dict[str, str]) -> str:
    parts = [f"*{course} | {option['label']}*", answer.strip()]
    return "\n\n".join(part for part in parts if part)


def _validated_setting(value: str, allowed: set[str], fallback: str, *, upper: bool = False) -> str:
    candidate = value.upper() if upper else value
    return candidate if candidate in allowed else fallback


def build_chat_request(
    sender: str,
    text: str,
    mode_override: str | None = None,
    course_override: str | None = None,
) -> ChatRequest:
    settings = get_settings()
    course = course_override or _validated_setting(settings.whatsapp_default_course, COURSES, "CMA", upper=True)
    mode = mode_override or _validated_setting(settings.whatsapp_default_mode, MODES, "teach")
    level = _validated_setting(settings.whatsapp_default_level, LEVELS, "beginner")

    stripped = text.strip()
    if not course_override:
        upper_text = stripped.upper()
        for candidate_course in COURSES:
            for prefix in (f"{candidate_course}:", f"[{candidate_course}]"):
                if upper_text.startswith(prefix):
                    course = candidate_course
                    stripped = stripped[len(prefix) :].strip()
                    break
            else:
                continue
            break

    return ChatRequest(
        student_id=f"whatsapp:{sender}",
        course=course,
        message=stripped or text,
        level=level,
        mode=mode,
        use_cache=True,
    )


class WhatsAppClient:
    def __init__(self) -> None:
        self.settings = get_settings()

    @property
    def configured(self) -> bool:
        return bool(self.settings.whatsapp_use_mock or (self._token() and self.settings.whatsapp_phone_number_id))

    def _token(self) -> str:
        return self.settings.whatsapp_access_token or self.settings.whatsapp_token

    def _graph_base(self) -> str:
        if self.settings.whatsapp_graph_base:
            return self.settings.whatsapp_graph_base.rstrip("/")
        return f"https://graph.facebook.com/{self.settings.whatsapp_graph_api_version}"

    def _send_url(self) -> str:
        return f"{self._graph_base()}/{self.settings.whatsapp_phone_number_id}/messages"

    async def download_media(self, media_id: str) -> tuple[bytes, str]:
        if not self._token():
            raise RuntimeError("WhatsApp access token is required to download media")

        headers = {"Authorization": f"Bearer {self._token()}"}
        async with httpx.AsyncClient(timeout=self.settings.request_timeout_seconds) as client:
            metadata_response = await client.get(f"{self._graph_base()}/{media_id}", headers=headers)
            metadata_response.raise_for_status()
            metadata = metadata_response.json()
            media_url = str(metadata.get("url") or "")
            mime_type = str(metadata.get("mime_type") or "application/octet-stream").lower()
            declared_size = int(metadata.get("file_size") or 0)

            if not media_url:
                raise RuntimeError("WhatsApp media URL is missing")
            if not mime_type.startswith("image/"):
                raise UnsupportedFeedbackMediaError("Only image screenshots are supported")
            if declared_size > self.settings.whatsapp_max_media_bytes:
                raise UnsupportedFeedbackMediaError("WhatsApp image exceeds the configured size limit")

            media_response = await client.get(media_url, headers=headers)
            media_response.raise_for_status()
            image_bytes = media_response.content

        if not image_bytes:
            raise UnsupportedFeedbackMediaError("WhatsApp image download was empty")
        if len(image_bytes) > self.settings.whatsapp_max_media_bytes:
            raise UnsupportedFeedbackMediaError("WhatsApp image exceeds the configured size limit")
        return image_bytes, mime_type

    async def _post_message(self, payload: dict) -> dict[str, Any]:
        # Read the current runtime value at the final outbound boundary so the
        # dashboard kill switch also fences workers that were already running.
        if not get_settings().whatsapp_messaging_enabled:
            raise RuntimeError("WhatsApp messaging is paused from the admin dashboard")
        if self.settings.whatsapp_use_mock:
            logger.info("[Mock WhatsApp] %s", payload)
            return {"mock": True, "payload": payload}

        if not self.configured:
            raise RuntimeError("WhatsApp access token and phone number ID are required")

        headers = {
            "Authorization": f"Bearer {self._token()}",
            "Content-Type": "application/json",
        }

        async with httpx.AsyncClient(timeout=self.settings.request_timeout_seconds) as client:
            response = await client.post(self._send_url(), headers=headers, json=payload)
            if response.is_error:
                try:
                    error = (response.json() or {}).get("error") or {}
                except ValueError:
                    error = {}
                code = error.get("code")
                error_type = error.get("type")
                message = error.get("message") or response.reason_phrase
                logger.error(
                    "WhatsApp send rejected status=%s code=%s type=%s message=%s",
                    response.status_code,
                    code,
                    error_type,
                    message,
                )
                raise RuntimeError(
                    f"Meta rejected the WhatsApp reply (HTTP {response.status_code}, code {code}): {message}"
                )
            try:
                result = response.json()
            except ValueError:
                result = {"raw": response.text}

        logger.info(
            "WhatsApp send accepted type=%s to=%s response=%s",
            payload.get("type"),
            payload.get("to"),
            result,
        )
        return result

    async def send_text(self, to: str, body: str) -> dict[str, Any]:
        to = normalize_wa_id(to)
        payload = {
            "messaging_product": "whatsapp",
            "recipient_type": "individual",
            "to": to,
            "type": "text",
            "text": {
                "preview_url": False,
                "body": body,
            },
        }
        return await self._post_message(payload)

    async def send_template(
        self,
        to: str,
        template_name: str,
        language_code: str,
        components: list[dict[str, Any]] | None = None,
    ) -> dict[str, Any]:
        to = normalize_wa_id(to)
        template: dict[str, Any] = {
            "name": template_name,
            "language": {"code": language_code},
        }
        if components:
            template["components"] = components

        payload = {
            "messaging_product": "whatsapp",
            "recipient_type": "individual",
            "to": to,
            "type": "template",
            "template": template,
        }
        return await self._post_message(payload)

    async def send_start_message(
        self,
        to: str,
        template_name: str | None = None,
        language_code: str | None = None,
    ) -> dict[str, Any]:
        return await self.send_template(
            to=to,
            template_name=template_name or self.settings.whatsapp_start_template_name,
            language_code=language_code or self.settings.whatsapp_start_template_language,
        )

    async def send_program_menu(self, to: str, courses: set[str] | None = None) -> dict[str, Any]:
        to = normalize_wa_id(to)
        available = set(courses or COURSES)
        rows = [
            {
                "id": option["id"],
                "title": option["label"],
                "description": option["description"],
            }
            for option in PROGRAM_OPTIONS
            if option["course"] in available
        ]
        if not rows:
            raise RuntimeError("At least one enrolled course is required")
        payload = {
            "messaging_product": "whatsapp",
            "recipient_type": "individual",
            "to": to,
            "type": "interactive",
            "interactive": {
                "type": "list",
                "header": {"type": "text", "text": "NorthStar Academy"},
                "body": {"text": "Choose one of your enrolled courses."},
                "footer": {"text": "Only enrolled courses are shown."},
                "action": {
                    "button": "Select course",
                    "sections": [{"title": "Your courses", "rows": rows}],
                },
            },
        }
        return await self._post_message(payload)

    async def send_mode_menu(self, to: str, course: str) -> dict[str, Any]:
        to = normalize_wa_id(to)
        rows = [
            {
                "id": option["id"],
                "title": option["label"],
                "description": option["description"],
            }
            for option in MODE_OPTIONS
        ]
        payload = {
            "messaging_product": "whatsapp",
            "recipient_type": "individual",
            "to": to,
            "type": "interactive",
            "interactive": {
                "type": "list",
                "header": {"type": "text", "text": course},
                "body": {"text": "Choose how I should help you."},
                "footer": {"text": "Reply 1-5. Type feedback to report an issue."},
                "action": {
                    "button": "Select option",
                    "sections": [{"title": "Learning modes", "rows": rows}],
                },
            },
        }
        return await self._post_message(payload)


class WhatsAppBot:
    def __init__(self) -> None:
        self.cache = Cache()
        self.client = WhatsAppClient()
        self.vision = GeminiService()
        self.mentor_factory = MentorService
        self.feedback_store_factory = AdminStore
        self.settings = get_settings()

    async def _save_feedback(
        self,
        message: WhatsAppIncomingMessage,
        description: str,
        *,
        include_image: bool,
        store: AdminStore | None = None,
    ) -> None:
        image_bytes = None
        mime_type = None
        if include_image:
            image_bytes, mime_type = await self.client.download_media(message.media_id or "")
        feedback_store = store or self.feedback_store_factory()
        feedback = await feedback_store.create_feedback(
            whatsapp_message_id=message.message_id,
            sender=normalize_wa_id(message.sender),
            profile_name=message.profile_name,
            message=description,
            image_bytes=image_bytes,
            mime_type=mime_type,
        )
        await self.client.send_text(
            message.sender,
            "Thank you. Your feedback was sent to the NorthStar Academy team. "
            f"Reference: *{feedback['id'].split('-', 1)[0].upper()}*",
        )
        # Keep the durable session until acknowledgement succeeds so Meta retries
        # cannot fall through into the tutoring flow after a transient send error.
        await feedback_store.clear_feedback_session(normalize_wa_id(message.sender))

    async def _handle_feedback_message(
        self,
        message: WhatsAppIncomingMessage,
        *,
        is_image: bool,
        incoming: str,
    ) -> bool:
        marker, marked_description = split_feedback_marker(incoming)
        feedback_store = self.feedback_store_factory()
        normalized_sender = normalize_wa_id(message.sender)
        feedback_state = await feedback_store.get_feedback_session(normalized_sender)

        if feedback_state and not is_image and is_hi_trigger(incoming):
            # Menu commands always escape a pending feedback conversation so a
            # stale session can never swallow a student's next greeting.
            await feedback_store.clear_feedback_session(normalized_sender)
            feedback_state = None

        if feedback_state and incoming.casefold() in FEEDBACK_CANCEL_TRIGGERS and not is_image:
            await self.client.send_text(message.sender, "Feedback cancelled. Type *menu* to continue learning.")
            await feedback_store.clear_feedback_session(normalized_sender)
            return True

        if marker:
            _, placeholder = split_feedback_marker(self.settings.whatsapp_feedback_prefill)
            if " ".join(marked_description.split()).casefold() == " ".join(placeholder.split()).casefold():
                marked_description = ""
            if is_image:
                try:
                    await self._save_feedback(
                        message,
                        marked_description or "Screenshot feedback",
                        include_image=True,
                        store=feedback_store,
                    )
                except (UnsupportedFeedbackMediaError, ValueError):
                    logger.exception("Could not save WhatsApp feedback screenshot message=%s", message.message_id)
                    await feedback_store.set_feedback_session(
                        normalized_sender,
                        "awaiting_submission",
                        self.settings.whatsapp_feedback_session_ttl_seconds,
                    )
                    await self.client.send_text(
                        message.sender,
                        "I could not save that screenshot. Please send a JPG or PNG under "
                        f"{max(1, self.settings.whatsapp_max_media_bytes // (1024 * 1024))} MB, or type your feedback.",
                    )
                    return True
                except Exception:
                    logger.exception("Transient WhatsApp feedback screenshot failure message=%s", message.message_id)
                    raise
                return True
            if marked_description:
                try:
                    await self._save_feedback(
                        message,
                        marked_description,
                        include_image=False,
                        store=feedback_store,
                    )
                except ValueError:
                    logger.exception("Could not save WhatsApp text feedback message=%s", message.message_id)
                    await self.client.send_text(
                        message.sender,
                        "I could not accept that feedback because its daily or storage limit was reached.",
                    )
                    return True
                except Exception:
                    logger.exception("Transient WhatsApp text feedback failure message=%s", message.message_id)
                    raise
                return True
            await feedback_store.set_feedback_session(
                normalized_sender,
                "awaiting_submission",
                self.settings.whatsapp_feedback_session_ttl_seconds,
            )
            await self.client.send_text(
                message.sender,
                "Please describe the change or error in one message. To include a screenshot, attach a JPG or PNG "
                "and write the description in its caption. Type *cancel* to exit.",
            )
            return True

        if not feedback_state:
            return False

        if is_image:
            try:
                await self._save_feedback(
                    message,
                    incoming or "Screenshot feedback",
                    include_image=True,
                    store=feedback_store,
                )
            except (UnsupportedFeedbackMediaError, ValueError):
                logger.exception("Could not save WhatsApp feedback screenshot message=%s", message.message_id)
                await self.client.send_text(
                    message.sender,
                    "I could not save that screenshot. Please send a JPG or PNG under "
                    f"{max(1, self.settings.whatsapp_max_media_bytes // (1024 * 1024))} MB, or type your feedback.",
                )
                return True
            except Exception:
                logger.exception("Transient WhatsApp feedback screenshot failure message=%s", message.message_id)
                raise
            return True
        if incoming:
            try:
                await self._save_feedback(
                    message,
                    incoming,
                    include_image=False,
                    store=feedback_store,
                )
            except ValueError:
                logger.exception("Could not save WhatsApp text feedback message=%s", message.message_id)
                await self.client.send_text(
                    message.sender,
                    "I could not accept that feedback because its daily or storage limit was reached.",
                )
                return True
            except Exception:
                logger.exception("Transient WhatsApp text feedback failure message=%s", message.message_id)
                raise
            return True
        await self.client.send_text(message.sender, "Please type your feedback or attach a JPG or PNG screenshot.")
        return True

    async def handle_payload(self, payload: dict[str, Any]) -> None:
        failed_messages: list[str] = []
        for event in extract_status_events(payload):
            if event.errors:
                logger.warning(
                    "WhatsApp status %s for message=%s recipient=%s errors=%s",
                    event.status,
                    event.message_id,
                    event.recipient_id,
                    event.errors,
                )
            else:
                logger.info(
                    "WhatsApp status %s for message=%s recipient=%s conversation=%s",
                    event.status,
                    event.message_id,
                    event.recipient_id,
                    event.conversation_id,
                )

        for message in extract_incoming_messages(payload):
            configured_phone_id = self.settings.whatsapp_phone_number_id.strip()
            if configured_phone_id and message.phone_number_id != configured_phone_id:
                # A single Meta/Xolox callback can receive events for several
                # business phone numbers. Never let this bot answer an event
                # addressed to another number. Meta includes phone_number_id on
                # every genuine Cloud API message event; a missing ID is accepted
                # only by explicit mock/test configurations.
                if message.phone_number_id or not (
                    self.settings.whatsapp_use_mock
                    or self.settings.app_environment == "test"
                ):
                    logger.warning(
                        "WhatsApp inbound ignored for non-configured phone message=%s target_phone_id=%s configured_phone_id=%s",
                        message.message_id,
                        message.phone_number_id or "missing",
                        configured_phone_id,
                    )
                    continue
            logger.info(
                "WhatsApp inbound message=%s sender=%s target_phone_id=%s type=%s has_media=%s",
                message.message_id,
                message.sender,
                message.phone_number_id or "missing",
                message.message_type,
                bool(message.media_id),
            )
            dedupe_key = f"whatsapp:inbound:done:{message.message_id}"
            processing_key = f"whatsapp:inbound:processing:{message.message_id}"
            if await self.cache.get_text(dedupe_key):
                logger.info("WhatsApp duplicate skipped message=%s", message.message_id)
                continue
            if not await self.cache.set_if_absent(
                processing_key,
                ttl_seconds=self.settings.whatsapp_inbound_processing_ttl_seconds,
            ):
                logger.info("WhatsApp duplicate already processing message=%s", message.message_id)
                # Do not report success to the durable queue: a previous worker may
                # have died before clearing this marker. Retry after its TTL instead.
                raise WhatsAppInboundBusyError(
                    message.message_id,
                    self.settings.whatsapp_inbound_processing_ttl_seconds,
                )

            try:
                await self._handle_message(message)
            except asyncio.CancelledError:
                # A graceful shutdown/recovery must not leave a Redis marker that
                # makes the durable queue's next attempt look successfully handled.
                await asyncio.shield(self.cache.delete(processing_key))
                raise
            except Exception:
                logger.exception("Failed to process WhatsApp message %s", message.message_id)
                # The durable webhook queue retries transient media, storage, or provider failures.
                await self.cache.delete(processing_key)
                failed_messages.append(message.message_id)
            else:
                await self.cache.set_text(dedupe_key, "1", ttl_seconds=86400)
                await self.cache.delete(processing_key)
        if failed_messages:
            raise RuntimeError(
                "WhatsApp processing failed for message(s): " + ", ".join(failed_messages)
            )

    async def _handle_message(self, message: WhatsAppIncomingMessage) -> None:
        is_image = message.message_type == "image" and bool(message.media_id)
        incoming = message.text.strip() if message.text else ""
        if await self._handle_feedback_message(message, is_image=is_image, incoming=incoming):
            return

        open_cma_access = self.settings.whatsapp_open_cma_access
        if not is_image and not message.text:
            await self.client.send_text(
                message.sender,
                "Please send a text question, an MCQ screenshot, or choose one menu option.",
            )
            return

        incoming_trigger = normalized_hi_trigger(incoming)
        is_hi_message = not is_image and is_hi_trigger(incoming)
        wants_course_list = is_hi_message and incoming_trigger in {"program", "programs"}
        selected_course = await self.cache.get_text(program_session_key(message.sender))
        requested_program = (
            None
            if is_image or incoming.strip().isdigit()
            else normalize_program_selection(incoming)[0]
        )
        needs_enrollment_lookup = (
            not open_cma_access
            or wants_course_list
            or (
                not is_hi_message
                and requested_program is None
                and selected_course not in {None, "CMA"}
            )
            or (requested_program is not None and requested_program["course"] != "CMA")
        )
        enrolled_courses = (
            await get_enrolled_courses(self.cache, message.sender)
            if needs_enrollment_lookup
            else set()
        )
        if open_cma_access:
            # Open CMA access is an effective permission only. Do not persist
            # unknown phone numbers or overwrite explicit course assignments.
            enrolled_courses = set(enrolled_courses)
            enrolled_courses.add("CMA")
        if not enrolled_courses:
            await self.client.send_text(
                message.sender,
                "This NorthStar Academy mentor is available only to enrolled students. Please contact NorthStar Academy support.",
            )
            return

        if selected_course not in enrolled_courses:
            selected_course = None
        if selected_course is None and (open_cma_access or len(enrolled_courses) == 1):
            selected_course = "CMA" if open_cma_access else next(iter(enrolled_courses))
            await self.cache.set_text(
                program_session_key(message.sender),
                selected_course,
                ttl_seconds=self.settings.whatsapp_session_ttl_seconds,
            )

        if is_hi_message:
            await self.cache.delete(mode_session_key(message.sender))
            await self.cache.set_text(
                conversation_started_key(message.sender),
                "1",
                ttl_seconds=self.settings.whatsapp_session_ttl_seconds,
            )
            if open_cma_access and wants_course_list and len(enrolled_courses) > 1:
                await self.cache.delete(program_session_key(message.sender))
                await self.client.send_text(
                    message.sender,
                    "Choose one of your available courses.\n\n" + program_menu(enrolled_courses),
                )
                await self.client.send_program_menu(message.sender, enrolled_courses)
            elif open_cma_access:
                selected_course = "CMA"
                await self.cache.set_text(
                    program_session_key(message.sender),
                    selected_course,
                    ttl_seconds=self.settings.whatsapp_session_ttl_seconds,
                )
                await self.cache.set_text(
                    mode_session_key(message.sender),
                    "teach",
                    ttl_seconds=self.settings.whatsapp_session_ttl_seconds,
                )
                await self.client.send_text(
                    message.sender,
                    "Welcome to NorthStar Academy. *CMA | Teach* is ready.\n\n"
                    "Ask any CMA question now, or upload a clear image of a CMA question. "
                    "Teach is already selected; the menu below lets you change learning mode if needed.\n\n"
                    + mode_menu(selected_course),
                )
                await self.client.send_mode_menu(message.sender, selected_course)
            elif len(enrolled_courses) > 1:
                await self.cache.delete(program_session_key(message.sender))
                await self.client.send_text(
                    message.sender,
                    "Welcome to NorthStar Academy. Choose one of your enrolled courses.\n\n"
                    + program_menu(enrolled_courses),
                )
                await self.client.send_program_menu(message.sender, enrolled_courses)
            else:
                selected_course = next(iter(enrolled_courses))
                await self.cache.set_text(
                    program_session_key(message.sender),
                    selected_course,
                    ttl_seconds=self.settings.whatsapp_session_ttl_seconds,
                )
                await self.client.send_text(
                    message.sender,
                    f"Welcome to NorthStar Academy. Your enrolled course is {selected_course}.\n\n"
                    + mode_menu(selected_course),
                )
                await self.client.send_mode_menu(message.sender, selected_course)
            return

        selected_program, program_inline_question = (None, "") if is_image else normalize_program_selection(incoming)
        if selected_course is not None and incoming.strip().isdigit():
            selected_program, program_inline_question = None, ""
        if selected_program:
            requested_course = selected_program["course"]
            if requested_course not in enrolled_courses:
                await self.client.send_text(
                    message.sender,
                    "That course is not enabled for your number. Your enrolled courses are: "
                    + ", ".join(ordered_courses(enrolled_courses)),
                )
                await self.client.send_program_menu(message.sender, enrolled_courses)
                return
            await self.cache.set_text(
                program_session_key(message.sender),
                requested_course,
                ttl_seconds=self.settings.whatsapp_session_ttl_seconds,
            )
            selected_course = requested_course
            if open_cma_access and requested_course == "CMA" and program_inline_question:
                await self.cache.set_text(
                    mode_session_key(message.sender),
                    "teach",
                    ttl_seconds=self.settings.whatsapp_session_ttl_seconds,
                )
                incoming = program_inline_question
            else:
                await self.cache.delete(mode_session_key(message.sender))
                await self.client.send_text(message.sender, mode_menu(requested_course))
                await self.client.send_mode_menu(message.sender, requested_course)
                if program_inline_question:
                    await self.client.send_text(
                        message.sender,
                        "I saved your course. Please choose Teach, Doubt Solving, Quiz, Revision, or Job Hunt before I answer that question.",
                    )
                return

        if selected_course is None:
            await self.client.send_text(
                message.sender,
                "Choose the enrolled course for this question first.\n\n" + program_menu(enrolled_courses),
            )
            await self.client.send_program_menu(message.sender, enrolled_courses)
            return

        if is_image:
            try:
                await self.client.send_text(
                    message.sender,
                    f"I received your {selected_course} image. I am reading it now and will send the explanation shortly.",
                )
                image_bytes, mime_type = await self.client.download_media(message.media_id or "")
                extracted = await self.vision.extract_question_from_image(
                    image_bytes=image_bytes,
                    mime_type=mime_type,
                    caption=message.text,
                    course=selected_course,
                )
                incoming = extracted[:5800]
                logger.info(
                    "WhatsApp image extracted message=%s sender=%s course=%s mime=%s bytes=%s chars=%s",
                    message.message_id,
                    message.sender,
                    selected_course,
                    mime_type,
                    len(image_bytes),
                    len(incoming),
                )
            except Exception:
                logger.exception("Failed to process WhatsApp image %s", message.media_id)
                await self.client.send_text(
                    message.sender,
                    "I could not read that screenshot. Please send a clear JPG or PNG image, or type the course question.",
                )
                return

        selected_option, inline_question = (None, "") if is_image else normalize_mode_selection(incoming)
        if selected_option:
            await self.cache.set_text(
                mode_session_key(message.sender),
                selected_option["mode"],
                ttl_seconds=self.settings.whatsapp_session_ttl_seconds,
            )
            if not inline_question:
                await self.client.send_text(message.sender, mode_prompt(selected_option))
                return
            incoming = inline_question

        selected_mode = await self.cache.get_text(mode_session_key(message.sender))
        if selected_mode not in MODES:
            selected_mode = (
                "teach"
                if open_cma_access and selected_course == "CMA"
                else _validated_setting(
                    self.settings.whatsapp_default_mode,
                    MODES,
                    "teach",
                )
            )
            await self.cache.set_text(
                mode_session_key(message.sender),
                selected_mode,
                ttl_seconds=self.settings.whatsapp_session_ttl_seconds,
            )

        selected_option = next(option for option in MODE_OPTIONS if option["mode"] == selected_mode)

        mentor = None
        usage_started = time.perf_counter()
        usage_status = "success"
        try:
            mentor = self.mentor_factory()
            chat_request = build_chat_request(
                message.sender,
                incoming,
                mode_override=selected_mode,
                course_override=selected_course,
            )
            if is_image and hasattr(mentor, "answer_image"):
                response = await mentor.answer_image(chat_request)
            else:
                response = await mentor.answer(chat_request)
            answer = format_whatsapp_answer(response.answer, selected_course, selected_option)
            logger.info(
                "WhatsApp mentor answer ready message=%s sender=%s course=%s chars=%s refused=%s",
                message.message_id,
                message.sender,
                selected_course,
                len(response.answer),
                response.answer
                in {
                    MentorService.refusal(),
                    MentorService.image_refusal(selected_course),
                    MentorService.course_refusal(selected_course),
                },
            )
        except Exception:
            usage_status = "error"
            logger.exception("Mentor answer failed for WhatsApp sender %s", message.sender)
            answer = "I had trouble generating that answer right now. Please try again in a moment."
        finally:
            closer = getattr(mentor, "aclose", None)
            if closer is not None:
                try:
                    await closer()
                except Exception:
                    logger.exception("Could not close WhatsApp mentor resources")
            try:
                await AdminStore().record_usage(
                    student_id=f"whatsapp:{message.sender}",
                    channel="whatsapp",
                    course=selected_course,
                    mode=selected_mode,
                    status=usage_status,
                    latency_ms=round((time.perf_counter() - usage_started) * 1000),
                )
            except Exception:
                logger.exception("Could not record WhatsApp usage analytics")

        for chunk in split_whatsapp_text(answer, self.settings.whatsapp_max_reply_chars):
            await self.client.send_text(message.sender, chunk)


async def process_whatsapp_webhook(payload: dict[str, Any]) -> None:
    bot = WhatsAppBot()
    try:
        await bot.handle_payload(payload)
    finally:
        try:
            await bot.cache.aclose()
        except Exception:
            logger.exception("Could not close WhatsApp cache client")

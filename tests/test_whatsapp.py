import asyncio

import pytest
from openpyxl import Workbook, load_workbook

import app.whatsapp as whatsapp_module
from app.config import get_settings
from app.schemas import ChatResponse
from app.whatsapp import (
    WhatsAppBot,
    WhatsAppIncomingMessage,
    WhatsAppClient,
    build_chat_request,
    extract_incoming_messages,
    extract_status_events,
    normalize_mode_selection,
    normalize_program_selection,
    mode_menu,
    program_menu,
    format_whatsapp_answer,
    get_enrolled_course,
    get_enrolled_courses,
    is_hi_trigger,
    remove_enrolled_course,
    set_enrolled_course,
    split_whatsapp_text,
    text_menu,
)
from scripts.enrollments.create_enrollment_workbook import create_workbook


def create_enrollment_test_workbook(path, rows=()):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Enrollments"
    sheet.append(["phone_number", "course", "active", "student_name", "notes"])
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    workbook.close()


@pytest.fixture(autouse=True)
def isolate_enrollment_sources(monkeypatch, tmp_path):
    monkeypatch.setenv("RUNTIME_CONFIG_FILE", str(tmp_path / "runtime-config.json"))
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_FILE", "")
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_GOOGLE_SHEET_ID", "")
    monkeypatch.setenv("WHATSAPP_OPEN_CMA_ACCESS", "false")
    get_settings.cache_clear()
    whatsapp_module._load_enrollment_workbook.cache_clear()
    yield
    get_settings.cache_clear()
    whatsapp_module._load_enrollment_workbook.cache_clear()


class FakeWhatsAppCache:
    def __init__(self) -> None:
        self.values = {}

    async def set_if_absent(self, key: str, ttl_seconds: int) -> bool:
        return True

    async def get_text(self, key: str) -> str | None:
        return self.values.get(key)

    async def set_text(self, key: str, value: str, ttl_seconds: int) -> None:
        self.values[key] = value

    async def delete(self, key: str) -> None:
        self.values.pop(key, None)


class FakeWhatsAppClient:
    def __init__(self) -> None:
        self.texts = []
        self.program_menus = []
        self.mode_menus = []
        self.downloads = []

    async def send_text(self, to: str, body: str) -> dict:
        self.texts.append((to, body))
        return {}

    async def send_program_menu(self, to: str, courses=None) -> dict:
        self.program_menus.append((to, set(courses or set())))
        return {}

    async def send_mode_menu(self, to: str, course: str) -> dict:
        self.mode_menus.append((to, course))
        return {}

    async def download_media(self, media_id: str):
        self.downloads.append(media_id)
        return b"image-bytes", "image/png"


class FakeVision:
    def __init__(self, extracted_text):
        self.extracted_text = extracted_text
        self.calls = []

    async def extract_question_from_image(self, **kwargs):
        self.calls.append(kwargs)
        return self.extracted_text


class FakeMentor:
    def __init__(self):
        self.requests = []

    async def answer(self, request):
        self.requests.append(request)
        return ChatResponse(answer="The correct answer is explained step by step.", sources=[])

    async def answer_image(self, request):
        return await self.answer(request)


def test_extract_incoming_text_message():
    payload = {
        "entry": [
            {
                "changes": [
                    {
                        "value": {
                            "metadata": {"phone_number_id": "146239241916262"},
                            "contacts": [{"wa_id": "919999999999", "profile": {"name": "Student"}}],
                            "messages": [
                                {
                                    "from": "919999999999",
                                    "id": "wamid.test",
                                    "type": "text",
                                    "text": {"body": "CMA: Explain variance analysis"},
                                }
                            ],
                        }
                    }
                ]
            }
        ]
    }

    messages = extract_incoming_messages(payload)

    assert len(messages) == 1
    assert messages[0].sender == "919999999999"
    assert messages[0].message_id == "wamid.test"
    assert messages[0].phone_number_id == "146239241916262"
    assert messages[0].text == "CMA: Explain variance analysis"
    assert messages[0].profile_name == "Student"


@pytest.mark.parametrize("text", ["Hi", "HI!", "hi 👋", "Hello...", "/start!"])
def test_hi_trigger_accepts_case_punctuation_and_trailing_emoji(text):
    assert is_hi_trigger(text) is True


@pytest.mark.parametrize("text", ["this", "hi there", "highlight", "hello teacher"])
def test_hi_trigger_does_not_consume_normal_questions(text):
    assert is_hi_trigger(text) is False


def test_bot_ignores_messages_addressed_to_another_business_number(monkeypatch):
    monkeypatch.setenv("APP_ENVIRONMENT", "test")
    monkeypatch.setenv("WHATSAPP_PHONE_NUMBER_ID", "146239241916262")
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS", "")
    monkeypatch.setenv("WHATSAPP_OPEN_CMA_ACCESS", "true")
    get_settings.cache_clear()
    bot = WhatsAppBot()
    bot.cache = FakeWhatsAppCache()
    bot.client = FakeWhatsAppClient()
    payload = {
        "entry": [
            {
                "changes": [
                    {
                        "value": {
                            "metadata": {"phone_number_id": "921055841100882"},
                            "contacts": [{"wa_id": "919535210826"}],
                            "messages": [
                                {
                                    "from": "919535210826",
                                    "id": "wamid.wrong-business-number",
                                    "type": "text",
                                    "text": {"body": "Hi"},
                                }
                            ],
                        }
                    }
                ]
            }
        ]
    }

    asyncio.run(bot.handle_payload(payload))

    assert bot.client.texts == []
    assert bot.client.program_menus == []
    assert bot.client.mode_menus == []


def test_extract_incoming_image_message_with_caption():
    payload = {
        "entry": [
            {
                "changes": [
                    {
                        "value": {
                            "messages": [
                                {
                                    "from": "919999999999",
                                    "id": "wamid.image",
                                    "type": "image",
                                    "image": {
                                        "id": "media-123",
                                        "mime_type": "image/png",
                                        "caption": "Please explain this MCQ",
                                    },
                                }
                            ]
                        }
                    }
                ]
            }
        ]
    }

    messages = extract_incoming_messages(payload)

    assert messages[0].message_type == "image"
    assert messages[0].media_id == "media-123"
    assert messages[0].media_mime_type == "image/png"
    assert messages[0].text == "Please explain this MCQ"


def test_build_chat_request_allows_course_prefix():
    request = build_chat_request(
        "919999999999",
        "What is materiality?",
        mode_override="teach",
        course_override="CPA",
    )

    assert request.student_id == "whatsapp:919999999999"
    assert request.course == "CPA"
    assert request.mode == "teach"
    assert request.message == "What is materiality?"


def test_split_whatsapp_text_keeps_chunks_below_limit():
    chunks = split_whatsapp_text("word " * 100, max_chars=80)

    assert len(chunks) > 1
    assert all(len(chunk) <= 80 for chunk in chunks)


def test_whatsapp_answer_does_not_expose_sources():
    option = {"label": "Doubt Solving", "mode": "doubt_solving"}

    result = format_whatsapp_answer("The answer with explanation.", "CMA", option)

    assert "The answer with explanation." in result
    assert "Sources" not in result
    assert "Type menu" not in result


def test_interactive_list_reply_extracts_mode_id():
    payload = {
        "entry": [
            {
                "changes": [
                    {
                        "value": {
                            "messages": [
                                {
                                    "from": "919999999999",
                                    "id": "wamid.mode",
                                    "type": "interactive",
                                    "interactive": {
                                        "type": "list_reply",
                                        "list_reply": {"id": "mode:job_hunt", "title": "Job Hunt"},
                                    },
                                }
                            ]
                        }
                    }
                ]
            }
        ]
    }

    messages = extract_incoming_messages(payload)

    assert messages[0].text == "mode:job_hunt"


def test_interactive_list_reply_extracts_program_id():
    payload = {
        "entry": [
            {
                "changes": [
                    {
                        "value": {
                            "messages": [
                                {
                                    "from": "919999999999",
                                    "id": "wamid.program",
                                    "type": "interactive",
                                    "interactive": {
                                        "type": "list_reply",
                                        "list_reply": {"id": "program:cpa", "title": "CPA"},
                                    },
                                }
                            ]
                        }
                    }
                ]
            }
        ]
    }

    messages = extract_incoming_messages(payload)

    assert messages[0].text == "program:cpa"


def test_extract_status_events_includes_errors():
    payload = {
        "entry": [
            {
                "changes": [
                    {
                        "value": {
                            "statuses": [
                                {
                                    "id": "wamid.status",
                                    "status": "failed",
                                    "timestamp": "1710000000",
                                    "recipient_id": "919999999999",
                                    "conversation": {"id": "conv.test"},
                                    "errors": [{"code": 131026, "title": "Message undeliverable"}],
                                }
                            ]
                        }
                    }
                ]
            }
        ]
    }

    events = extract_status_events(payload)

    assert len(events) == 1
    assert events[0].message_id == "wamid.status"
    assert events[0].status == "failed"
    assert events[0].recipient_id == "919999999999"
    assert events[0].errors[0]["code"] == 131026


def test_send_start_message_uses_whatsapp_template_payload(monkeypatch):
    monkeypatch.setenv("WHATSAPP_USE_MOCK", "true")
    monkeypatch.setenv("WHATSAPP_START_TEMPLATE_NAME", "hello_world")
    monkeypatch.setenv("WHATSAPP_START_TEMPLATE_LANGUAGE", "en_US")
    get_settings.cache_clear()

    client = WhatsAppClient()
    result = asyncio.run(client.send_start_message("9916039894"))
    get_settings.cache_clear()

    payload = result["payload"]
    assert payload["to"] == "919916039894"
    assert payload["type"] == "template"
    assert payload["template"]["name"] == "hello_world"
    assert payload["template"]["language"]["code"] == "en_US"


def test_enrolled_customer_hi_opens_only_enrolled_course_mode_menu(monkeypatch):
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS", "919916039894:CMA")
    get_settings.cache_clear()
    bot = WhatsAppBot()
    bot.cache = FakeWhatsAppCache()
    bot.client = FakeWhatsAppClient()

    message = WhatsAppIncomingMessage(
        message_id="wamid.reply",
        sender="919916039894",
        text="hi",
        message_type="text",
    )
    asyncio.run(bot._handle_message(message))
    get_settings.cache_clear()

    assert bot.client.program_menus == []
    assert bot.client.mode_menus == [("919916039894", "CMA")]


def test_phone_only_sheet_row_defaults_to_active_cma_access(tmp_path, monkeypatch):
    workbook_path = tmp_path / "phone-only-enrollment.xlsx"
    create_enrollment_test_workbook(
        workbook_path,
        [("919876543210", "", "", "New CMA student", "")],
    )
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_FILE", str(workbook_path))
    get_settings.cache_clear()
    whatsapp_module._load_enrollment_workbook.cache_clear()
    bot = WhatsAppBot()
    bot.cache = FakeWhatsAppCache()
    bot.client = FakeWhatsAppClient()

    asyncio.run(
        bot._handle_message(
            WhatsAppIncomingMessage(
                message_id="wamid.phone-only-hi",
                sender="919876543210",
                text="Hi",
                message_type="text",
            )
        )
    )

    assert bot.client.program_menus == []
    assert bot.client.mode_menus == [("919876543210", "CMA")]
    assert bot.cache.values[whatsapp_module.program_session_key("919876543210")] == "CMA"


def test_customer_message_then_mode_selection_then_question_gets_mentor_answer(monkeypatch):
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS", "919916039894:CMA")
    get_settings.cache_clear()
    bot = WhatsAppBot()
    bot.cache = FakeWhatsAppCache()
    bot.client = FakeWhatsAppClient()
    mentor = FakeMentor()
    bot.mentor_factory = lambda: mentor

    async def customer_journey():
        await bot._handle_message(
            WhatsAppIncomingMessage(
                message_id="wamid.journey-hi",
                sender="919916039894",
                text="Hi",
                message_type="text",
            )
        )
        await bot._handle_message(
            WhatsAppIncomingMessage(
                message_id="wamid.journey-mode",
                sender="919916039894",
                text="mode:teach",
                message_type="text",
            )
        )
        await bot._handle_message(
            WhatsAppIncomingMessage(
                message_id="wamid.journey-question",
                sender="919916039894",
                text="Explain standard costing",
                message_type="text",
            )
        )

    asyncio.run(customer_journey())
    get_settings.cache_clear()

    assert bot.client.mode_menus == [("919916039894", "CMA")]
    assert mentor.requests[0].course == "CMA"
    assert mentor.requests[0].mode == "teach"
    assert mentor.requests[0].message == "Explain standard costing"
    assert "correct answer is explained step by step" in bot.client.texts[-1][1]


def test_unknown_customer_gets_cma_teach_option_and_can_ask_a_question(monkeypatch):
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS", "919916039894:CMA")
    monkeypatch.setenv("WHATSAPP_OPEN_CMA_ACCESS", "true")
    get_settings.cache_clear()
    bot = WhatsAppBot()
    bot.cache = FakeWhatsAppCache()
    bot.client = FakeWhatsAppClient()
    mentor = FakeMentor()
    bot.mentor_factory = lambda: mentor

    async def customer_journey():
        await bot._handle_message(
            WhatsAppIncomingMessage(
                message_id="wamid.open-cma-hi",
                sender="919999999999",
                text="hi",
                message_type="text",
            )
        )
        await bot._handle_message(
            WhatsAppIncomingMessage(
                message_id="wamid.open-cma-teach",
                sender="919999999999",
                text="mode:teach",
                message_type="text",
            )
        )
        await bot._handle_message(
            WhatsAppIncomingMessage(
                message_id="wamid.open-cma-question",
                sender="919999999999",
                text="Explain standard costing",
                message_type="text",
            )
        )

    asyncio.run(customer_journey())
    get_settings.cache_clear()

    assert bot.client.program_menus == []
    assert bot.client.mode_menus == [("919999999999", "CMA")]
    assert "CMA" in bot.client.texts[0][1]
    assert "Teach" in bot.client.texts[0][1]
    assert mentor.requests[0].student_id == "whatsapp:919999999999"
    assert mentor.requests[0].course == "CMA"
    assert mentor.requests[0].mode == "teach"
    assert mentor.requests[0].message == "Explain standard costing"
    assert "correct answer is explained step by step" in bot.client.texts[-1][1]


def test_unknown_customer_can_ask_a_cma_question_as_the_first_message(monkeypatch):
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS", "")
    monkeypatch.setenv("WHATSAPP_OPEN_CMA_ACCESS", "true")
    get_settings.cache_clear()
    bot = WhatsAppBot()
    bot.cache = FakeWhatsAppCache()
    bot.client = FakeWhatsAppClient()
    mentor = FakeMentor()
    bot.mentor_factory = lambda: mentor

    asyncio.run(
        bot._handle_message(
            WhatsAppIncomingMessage(
                message_id="wamid.open-cma-direct-question",
                sender="919999999999",
                text="Explain responsibility accounting",
                message_type="text",
            )
        )
    )
    get_settings.cache_clear()

    assert bot.client.program_menus == []
    assert bot.client.mode_menus == []
    assert bot.cache.values[whatsapp_module.program_session_key("919999999999")] == "CMA"
    assert bot.cache.values[whatsapp_module.mode_session_key("919999999999")] == "teach"
    assert mentor.requests[0].course == "CMA"
    assert mentor.requests[0].mode == "teach"
    assert mentor.requests[0].message == "Explain responsibility accounting"
    assert "CMA | Teach" in bot.client.texts[-1][1]


def test_open_cma_does_not_wait_for_enrollment_lookup(monkeypatch):
    monkeypatch.setenv("WHATSAPP_OPEN_CMA_ACCESS", "true")
    get_settings.cache_clear()
    bot = WhatsAppBot()
    bot.cache = FakeWhatsAppCache()
    bot.client = FakeWhatsAppClient()
    mentor = FakeMentor()
    bot.mentor_factory = lambda: mentor

    async def unavailable_enrollment_source(*_args, **_kwargs):
        raise AssertionError("Public CMA must not query the enrollment source")

    monkeypatch.setattr(whatsapp_module, "get_enrolled_courses", unavailable_enrollment_source)

    async def customer_journey():
        await bot._handle_message(
            WhatsAppIncomingMessage(
                message_id="wamid.open-cma-no-enrollment-hi",
                sender="919999999999",
                text="Hi",
                message_type="text",
            )
        )
        await bot._handle_message(
            WhatsAppIncomingMessage(
                message_id="wamid.open-cma-no-enrollment-question",
                sender="919999999999",
                text="Explain transfer pricing",
                message_type="text",
            )
        )

    asyncio.run(customer_journey())
    get_settings.cache_clear()

    assert mentor.requests[0].course == "CMA"
    assert mentor.requests[0].mode == "teach"
    assert mentor.requests[0].message == "Explain transfer pricing"


def test_open_cma_inline_course_question_is_answered_immediately_in_teach(monkeypatch):
    monkeypatch.setenv("WHATSAPP_OPEN_CMA_ACCESS", "true")
    get_settings.cache_clear()
    bot = WhatsAppBot()
    bot.cache = FakeWhatsAppCache()
    bot.client = FakeWhatsAppClient()
    mentor = FakeMentor()
    bot.mentor_factory = lambda: mentor

    asyncio.run(
        bot._handle_message(
            WhatsAppIncomingMessage(
                message_id="wamid.open-cma-inline-question",
                sender="919999999999",
                text="CMA: Explain absorption costing",
                message_type="text",
            )
        )
    )
    get_settings.cache_clear()

    assert bot.client.mode_menus == []
    assert mentor.requests[0].course == "CMA"
    assert mentor.requests[0].mode == "teach"
    assert mentor.requests[0].message == "Explain absorption costing"
    assert "CMA | Teach" in bot.client.texts[-1][1]


def test_open_cma_programs_command_lists_additional_enrolled_courses(monkeypatch):
    monkeypatch.setenv("WHATSAPP_OPEN_CMA_ACCESS", "true")
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS", "919999999999:CFA")
    get_settings.cache_clear()
    bot = WhatsAppBot()
    bot.cache = FakeWhatsAppCache()
    bot.client = FakeWhatsAppClient()

    asyncio.run(
        bot._handle_message(
            WhatsAppIncomingMessage(
                message_id="wamid.open-cma-programs",
                sender="919999999999",
                text="programs",
                message_type="text",
            )
        )
    )
    get_settings.cache_clear()

    assert bot.client.program_menus == [("919999999999", {"CMA", "CFA"})]
    assert bot.client.mode_menus == []


def test_unknown_customer_is_refused_when_open_cma_access_is_disabled(monkeypatch):
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS", "919916039894:CMA")
    monkeypatch.setenv("WHATSAPP_OPEN_CMA_ACCESS", "false")
    get_settings.cache_clear()
    bot = WhatsAppBot()
    bot.cache = FakeWhatsAppCache()
    bot.client = FakeWhatsAppClient()

    asyncio.run(
        bot._handle_message(
            WhatsAppIncomingMessage(
                message_id="wamid.closed-cma-hi",
                sender="919999999999",
                text="hi",
                message_type="text",
            )
        )
    )
    get_settings.cache_clear()

    assert "only to enrolled students" in bot.client.texts[0][1]
    assert bot.client.program_menus == []
    assert bot.client.mode_menus == []


def test_unknown_customer_can_upload_a_cma_image_and_receive_a_teach_answer(monkeypatch):
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS", "919916039894:CMA")
    monkeypatch.setenv("WHATSAPP_OPEN_CMA_ACCESS", "true")
    get_settings.cache_clear()
    bot = WhatsAppBot()
    bot.cache = FakeWhatsAppCache()
    bot.client = FakeWhatsAppClient()
    bot.vision = FakeVision("Which variance is favorable? A. Price B. Quantity")
    mentor = FakeMentor()
    bot.mentor_factory = lambda: mentor

    async def image_journey():
        await bot._handle_message(
            WhatsAppIncomingMessage(
                message_id="wamid.open-cma-image",
                sender="919999999999",
                text="Explain every option",
                message_type="image",
                media_id="media-open-cma",
                media_mime_type="image/png",
            )
        )

    asyncio.run(image_journey())
    get_settings.cache_clear()

    assert bot.client.mode_menus == []
    assert bot.client.downloads == ["media-open-cma"]
    assert bot.cache.values[whatsapp_module.program_session_key("919999999999")] == "CMA"
    assert bot.cache.values[whatsapp_module.mode_session_key("919999999999")] == "teach"
    assert "I received your CMA image" in bot.client.texts[-2][1]
    assert bot.vision.calls == [
        {
            "image_bytes": b"image-bytes",
            "mime_type": "image/png",
            "caption": "Explain every option",
            "course": "CMA",
        }
    ]
    assert mentor.requests[0].student_id == "whatsapp:919999999999"
    assert mentor.requests[0].course == "CMA"
    assert mentor.requests[0].mode == "teach"
    assert "Which variance is favorable?" in mentor.requests[0].message
    assert "correct answer is explained step by step" in bot.client.texts[-1][1]


def test_enrolled_customer_cannot_switch_course(monkeypatch):
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS", "919916039894:CMA")
    get_settings.cache_clear()
    bot = WhatsAppBot()
    bot.cache = FakeWhatsAppCache()
    bot.client = FakeWhatsAppClient()

    message = WhatsAppIncomingMessage(
        message_id="wamid.switch",
        sender="919916039894",
        text="CPA",
        message_type="text",
    )
    asyncio.run(bot._handle_message(message))
    get_settings.cache_clear()

    assert "not enabled" in bot.client.texts[0][1]
    assert bot.client.program_menus == [("919916039894", {"CMA"})]
    assert bot.client.mode_menus == []


def test_dashboard_enrollment_revoke_overrides_static_fallback(monkeypatch):
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS", "919916039894:CMA;919916039894:CPA")
    get_settings.cache_clear()
    cache = FakeWhatsAppCache()

    remaining = asyncio.run(remove_enrolled_course(cache, "919916039894", "CMA"))
    assert remaining == {"CPA"}
    assert asyncio.run(get_enrolled_courses(cache, "919916039894")) == {"CPA"}
    get_settings.cache_clear()


def test_multi_course_customer_selects_only_an_enrolled_course(monkeypatch):
    monkeypatch.setenv(
        "WHATSAPP_ENROLLMENTS",
        "919916039894:CMA;919916039894:CPA;919916039894:CFA;919916039894:ACCA;919916039894:CS;919916039894:EA",
    )
    get_settings.cache_clear()
    bot = WhatsAppBot()
    bot.cache = FakeWhatsAppCache()
    bot.client = FakeWhatsAppClient()

    asyncio.run(
        bot._handle_message(
            WhatsAppIncomingMessage(
                message_id="wamid.multi-hi",
                sender="919916039894",
                text="hi",
                message_type="text",
            )
        )
    )

    assert bot.client.program_menus == [
        ("919916039894", {"CMA", "CPA", "CFA", "ACCA", "CS", "EA"})
    ]
    assert bot.client.mode_menus == []

    asyncio.run(
        bot._handle_message(
            WhatsAppIncomingMessage(
                message_id="wamid.select-cfa",
                sender="919916039894",
                text="CFA",
                message_type="text",
            )
        )
    )
    get_settings.cache_clear()

    assert bot.client.mode_menus == [("919916039894", "CFA")]
    assert bot.cache.values[whatsapp_module.program_session_key("919916039894")] == "CFA"


def test_enrolled_customer_image_is_extracted_and_answered_in_fixed_course(monkeypatch):
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS", "919916039894:CMA")
    get_settings.cache_clear()
    bot = WhatsAppBot()
    bot.cache = FakeWhatsAppCache()
    bot.client = FakeWhatsAppClient()
    bot.vision = FakeVision("Which variance is favorable? A. Price B. Quantity")
    mentor = FakeMentor()
    bot.mentor_factory = lambda: mentor

    message = WhatsAppIncomingMessage(
        message_id="wamid.image",
        sender="919916039894",
        text="Explain every option",
        message_type="image",
        media_id="media-123",
        media_mime_type="image/png",
    )
    asyncio.run(bot._handle_message(message))
    get_settings.cache_clear()

    assert bot.client.downloads == ["media-123"]
    assert "I received your CMA image" in bot.client.texts[0][1]
    assert bot.vision.calls[0]["course"] == "CMA"
    assert mentor.requests[0].course == "CMA"
    assert mentor.requests[0].mode == "teach"
    assert "Which variance is favorable?" in mentor.requests[0].message
    assert "correct answer is explained step by step" in bot.client.texts[-1][1]


def test_multi_course_image_uses_the_student_selected_course(monkeypatch):
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS", "919916039894:CMA;919916039894:CFA")
    get_settings.cache_clear()
    bot = WhatsAppBot()
    bot.cache = FakeWhatsAppCache()
    bot.cache.values[whatsapp_module.program_session_key("919916039894")] = "CFA"
    bot.client = FakeWhatsAppClient()
    bot.vision = FakeVision("Explain bond duration when yields increase.")
    mentor = FakeMentor()
    bot.mentor_factory = lambda: mentor

    asyncio.run(
        bot._handle_message(
            WhatsAppIncomingMessage(
                message_id="wamid.cfa-image",
                sender="919916039894",
                text="",
                message_type="image",
                media_id="media-cfa",
                media_mime_type="image/jpeg",
            )
        )
    )
    get_settings.cache_clear()

    assert bot.vision.calls[0]["course"] == "CFA"
    assert mentor.requests[0].course == "CFA"
    assert "CFA | Teach" in bot.client.texts[-1][1]


def test_normalize_mode_selection_accepts_numbers_names_and_inline_questions():
    option, question = normalize_mode_selection("2")
    assert option["mode"] == "doubt_solving"
    assert question == ""

    option, question = normalize_mode_selection("Job Hunt: prepare me for FP&A")
    assert option["mode"] == "job_hunt"
    assert question == "prepare me for FP&A"


def test_normalize_program_selection_accepts_numbers_names_and_inline_questions():
    option, question = normalize_program_selection("2")
    assert option["course"] == "CPA"
    assert question == ""

    option, question = normalize_program_selection("ACCA: Explain audit risk")
    assert option["course"] == "ACCA"
    assert question == "Explain audit risk"


def test_text_menu_starts_with_programs():
    menu = text_menu()

    assert menu == program_menu()
    assert "CMA" in menu
    assert "CPA" in menu
    assert "CFA" in menu
    assert "ACCA" in menu
    assert "CS" in menu
    assert "EA" in menu


def test_mode_menu_contains_only_allowed_modes():
    menu = mode_menu("CPA")

    assert "Teach" in menu
    assert "Doubt Solving" in menu
    assert "Quiz" in menu
    assert "Revision" in menu
    assert "Job Hunt" in menu


def test_google_sheet_is_authoritative_for_enrollment(monkeypatch):
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_GOOGLE_SHEET_ID", "sheet-id")
    monkeypatch.setenv("GOOGLE_SERVICE_ACCOUNT_FILE", "service-account.json")
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS", "919535210826:CPA")
    get_settings.cache_clear()
    whatsapp_module._google_sheet_cache = None
    whatsapp_module._google_sheet_cache_key = ""
    whatsapp_module._google_sheet_cache_expires_at = 0.0
    monkeypatch.setattr(
        whatsapp_module,
        "_fetch_google_sheet_enrollments",
        lambda *_args: {"919535210826": {"CMA", "CFA"}},
    )

    courses = asyncio.run(get_enrolled_courses(FakeWhatsAppCache(), "9535210826"))

    assert courses == {"CMA", "CFA"}
    get_settings.cache_clear()


def test_missing_google_sheet_row_denies_fallback_access(monkeypatch):
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_GOOGLE_SHEET_ID", "sheet-id")
    monkeypatch.setenv("GOOGLE_SERVICE_ACCOUNT_FILE", "service-account.json")
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS", "919535210826:CMA")
    get_settings.cache_clear()
    whatsapp_module._google_sheet_cache = None
    whatsapp_module._google_sheet_cache_key = ""
    whatsapp_module._google_sheet_cache_expires_at = 0.0
    monkeypatch.setattr(
        whatsapp_module,
        "_fetch_google_sheet_enrollments",
        lambda *_args: {"919999999999": {"CMA"}},
    )

    course = asyncio.run(get_enrolled_course(FakeWhatsAppCache(), "9535210826"))

    assert course is None
    get_settings.cache_clear()


def test_google_sheet_bulk_updates_use_one_writable_append(monkeypatch, tmp_path):
    captured = {}
    credentials = tmp_path / "service-account.json"
    credentials.write_text("{}", encoding="utf-8")
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_GOOGLE_SHEET_ID", "sheet-id")
    monkeypatch.setenv("GOOGLE_SERVICE_ACCOUNT_FILE", str(credentials))
    get_settings.cache_clear()

    def fake_append(spreadsheet_id, sheet_range, service_account_file, rows):
        captured.update(
            spreadsheet_id=spreadsheet_id,
            sheet_range=sheet_range,
            service_account_file=service_account_file,
            rows=rows,
        )

    async def fake_sheet():
        return {"447700900123": {"CMA"}}

    monkeypatch.setattr(whatsapp_module, "_append_google_sheet_rows_sync", fake_append)
    monkeypatch.setattr(whatsapp_module, "google_sheet_enrollments", fake_sheet)
    result = asyncio.run(
        whatsapp_module.bulk_set_enrollment_states(
            [("+44 7700 900123", "CMA", True, "Student", "Bulk import")]
        )
    )
    assert result == {"source": "google_sheet", "rows": 1, "granted": 1, "revoked": 0}
    assert captured["spreadsheet_id"] == "sheet-id"
    assert captured["rows"] == [("447700900123", "CMA", True, "Student", "Bulk import")]
    get_settings.cache_clear()


def test_enrollment_rows_disable_invalid_or_inactive_students():
    rows = [
        ["phone_number", "course", "active", "student_name", "notes"],
        ["9535210826", "CMA", "YES", "Student A", ""],
        ["9916039894", "CPA", "NO", "Student B", ""],
        ["9876543210", "UNKNOWN", "YES", "Student C", ""],
    ]

    enrollments = whatsapp_module._parse_enrollment_rows(rows)

    assert enrollments["919535210826"] == {"CMA"}
    assert enrollments["919916039894"] == set()
    assert enrollments["919876543210"] == set()


def test_enrollment_rows_accumulate_multiple_active_courses_and_last_row_wins():
    rows = [
        ["phone_number", "course", "active", "student_name", "notes"],
        ["9535210826", "CMA", "YES", "Student A", ""],
        ["9535210826", "CFA", "YES", "Student A", ""],
        ["9535210826", "CS", "YES", "Student A", ""],
        ["9535210826", "CMA", "NO", "Student A", "revoked"],
    ]

    enrollments = whatsapp_module._parse_enrollment_rows(rows)

    assert enrollments["919535210826"] == {"CFA", "CS"}


def test_dashboard_enrollment_changes_are_written_live_to_excel(tmp_path, monkeypatch):
    workbook_path = tmp_path / "whatsapp_enrollments.xlsx"
    create_enrollment_test_workbook(
        workbook_path,
        [
            ["919916039894", "CMA", "YES", "Student", ""],
            ["919916039894", "CPA", "NO", "Student", ""],
        ],
    )
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_FILE", str(workbook_path))
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_GOOGLE_SHEET_ID", "")
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS", "919916039894:CFA")
    get_settings.cache_clear()
    whatsapp_module._load_enrollment_workbook.cache_clear()
    cache = FakeWhatsAppCache()

    assert asyncio.run(set_enrolled_course(cache, "919916039894", "CPA")) == "CPA"
    assert asyncio.run(get_enrolled_courses(cache, "919916039894")) == {"CMA", "CPA"}
    assert asyncio.run(remove_enrolled_course(cache, "919916039894", "CMA")) == {"CPA"}

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    rows = list(workbook["Enrollments"].iter_rows(min_row=2, values_only=True))
    workbook.close()
    states = {(str(row[0]), str(row[1])): str(row[2]) for row in rows}
    assert states[("919916039894", "CMA")] == "NO"
    assert states[("919916039894", "CPA")] == "YES"
    assert asyncio.run(get_enrolled_courses(cache, "919916039894")) == {"CPA"}
    get_settings.cache_clear()


def test_legacy_redis_enrollment_is_migrated_to_excel_without_losing_courses(
    tmp_path,
    monkeypatch,
):
    workbook_path = tmp_path / "whatsapp_enrollments.xlsx"
    create_enrollment_test_workbook(workbook_path)
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_FILE", str(workbook_path))
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_GOOGLE_SHEET_ID", "")
    get_settings.cache_clear()
    whatsapp_module._load_enrollment_workbook.cache_clear()
    cache = FakeWhatsAppCache()
    key = whatsapp_module.enrollment_key("919916039894")
    cache.values[key] = "CMA,CPA"

    courses = asyncio.run(get_enrolled_courses(cache, "919916039894"))

    assert courses == {"CMA", "CPA"}
    assert key not in cache.values
    assert whatsapp_module.workbook_enrollments(str(workbook_path))["919916039894"] == {
        "CMA",
        "CPA",
    }
    get_settings.cache_clear()


def test_concurrent_excel_course_grants_preserve_both_updates(tmp_path, monkeypatch):
    workbook_path = tmp_path / "whatsapp_enrollments.xlsx"
    create_enrollment_test_workbook(workbook_path)
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_FILE", str(workbook_path))
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_GOOGLE_SHEET_ID", "")
    get_settings.cache_clear()
    whatsapp_module._load_enrollment_workbook.cache_clear()
    cache = FakeWhatsAppCache()

    async def grant_both():
        await asyncio.gather(
            set_enrolled_course(cache, "919876543210", "CMA"),
            set_enrolled_course(cache, "919876543210", "CPA"),
        )

    asyncio.run(grant_both())
    assert asyncio.run(get_enrolled_courses(cache, "919876543210")) == {"CMA", "CPA"}
    get_settings.cache_clear()


def test_excel_writer_ignores_preformatted_blank_template_rows(tmp_path, monkeypatch):
    workbook_path = tmp_path / "whatsapp_enrollments.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Enrollments"
    sheet.append(["phone_number", "course", "active", "student_name", "notes"])
    sheet.append(["919535210826", "CMA", "YES", "Existing student", ""])
    for row_number in range(3, 101):
        sheet.cell(row=row_number, column=1).number_format = "@"
    workbook.save(workbook_path)
    workbook.close()
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_FILE", str(workbook_path))
    get_settings.cache_clear()
    whatsapp_module._load_enrollment_workbook.cache_clear()

    asyncio.run(set_enrolled_course(FakeWhatsAppCache(), "919876543210", "CPA"))

    saved = load_workbook(workbook_path, read_only=False, data_only=True)
    sheet = saved["Enrollments"]
    assert sheet.cell(row=3, column=1).value == "919876543210"
    assert sheet.cell(row=3, column=2).value == "CPA"
    assert sheet.auto_filter.ref == "A1:E3"
    saved.close()
    get_settings.cache_clear()


def test_enrollment_workbook_generator_preserves_existing_files_and_samples_are_opt_in(tmp_path):
    workbook_path = tmp_path / "whatsapp_enrollments.xlsx"

    create_workbook(workbook_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    rows = [
        row
        for row in workbook["Enrollments"].iter_rows(min_row=2, values_only=True)
        if any(value not in (None, "") for value in row)
    ]
    workbook.close()
    assert rows == []

    with pytest.raises(FileExistsError, match="Refusing to overwrite"):
        create_workbook(workbook_path)

    create_workbook(workbook_path, with_sample=True, force=True)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sample = next(workbook["Enrollments"].iter_rows(min_row=2, values_only=True))
    workbook.close()
    assert sample[:3] == ("919999999999", "CMA", "YES")


def test_excel_replace_failure_keeps_original_workbook_and_raises_clear_error(
    tmp_path,
    monkeypatch,
):
    workbook_path = tmp_path / "whatsapp_enrollments.xlsx"
    create_enrollment_test_workbook(
        workbook_path,
        [["919876543210", "CMA", "YES", "Student", ""]],
    )
    original = workbook_path.read_bytes()
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_FILE", str(workbook_path))
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_GOOGLE_SHEET_ID", "")
    get_settings.cache_clear()
    whatsapp_module._load_enrollment_workbook.cache_clear()
    monkeypatch.setattr(
        whatsapp_module.os,
        "replace",
        lambda *_args: (_ for _ in ()).throw(PermissionError("locked")),
    )

    with pytest.raises(whatsapp_module.EnrollmentWorkbookError, match="open or not writable"):
        asyncio.run(set_enrolled_course(FakeWhatsAppCache(), "919876543210", "CPA"))

    assert workbook_path.read_bytes() == original
    assert list(tmp_path.glob(".whatsapp_enrollments.*.xlsx")) == []
    get_settings.cache_clear()

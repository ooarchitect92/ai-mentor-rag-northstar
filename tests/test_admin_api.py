import asyncio

import hashlib
import hmac
import json

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest

import app.main as main_module
import app.admin_api as admin_api_module
import app.whatsapp as whatsapp_module
from app.admin_store import AdminStore
from app.config import get_settings
from app.main import app, validate_startup_configuration
from app.schemas import ChatResponse
from scripts.enrollments.set_student_courses import normalize_phone, set_student_courses


@pytest.fixture
def admin_client(tmp_path, monkeypatch):
    monkeypatch.setenv("APP_ENVIRONMENT", "test")
    monkeypatch.setenv("ADMIN_TOKEN", "test-admin-token")
    monkeypatch.setenv("ADMIN_DATABASE_FILE", str(tmp_path / "admin.sqlite3"))
    monkeypatch.setenv("FEEDBACK_MEDIA_DIRECTORY", str(tmp_path / "feedback-media"))
    monkeypatch.setenv("RUNTIME_CONFIG_FILE", str(tmp_path / "runtime-config.json"))
    monkeypatch.setenv("MENTOR_SYSTEM_PROMPT_FILE", str(tmp_path / "mentor-prompt.txt"))
    monkeypatch.setenv("WHATSAPP_FEEDBACK_NUMBER", "919876543210")
    monkeypatch.setenv("WHATSAPP_ACCESS_TOKEN", "test-token")
    monkeypatch.setenv("WHATSAPP_PHONE_NUMBER_ID", "test-phone-id")
    monkeypatch.setenv("WHATSAPP_BUSINESS_ACCOUNT_ID", "test-business-id")
    monkeypatch.setenv("WHATSAPP_WEBHOOK_CALLBACK_URL", "https://example.test/webhooks/whatsapp")
    monkeypatch.setenv("NORTHSTAR_PUBLIC_BASE_URL", "https://mentor.example.test")
    monkeypatch.setenv("WHATSAPP_APP_SECRET", "test-app-secret")
    monkeypatch.setenv("WHATSAPP_RELAY_TOKEN", "test-relay-token-that-is-longer-than-32-characters")
    monkeypatch.setenv("WHATSAPP_OPEN_CMA_ACCESS", "true")
    monkeypatch.setenv("NVIDIA_API_KEY", "test-nvidia-key")
    monkeypatch.setenv("GEMINI_API_KEY", "test-gemini-key")
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS", "")
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_FILE", "")
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_GOOGLE_SHEET_ID", "")
    get_settings.cache_clear()
    whatsapp_module._load_enrollment_workbook.cache_clear()
    client = TestClient(app)
    yield client, {"x-admin-token": "test-admin-token"}
    get_settings.cache_clear()
    whatsapp_module._load_enrollment_workbook.cache_clear()


def test_admin_auth_static_dashboard_and_public_feedback_config(admin_client):
    client, headers = admin_client

    health = client.get("/health")
    assert health.status_code == 200
    assert health.json()["status"] == "ok"
    assert client.get("/live").json() == {"status": "ok"}

    assert client.get("/v1/admin/overview").status_code == 401
    overview = client.get("/v1/admin/overview", headers=headers)
    assert overview.status_code == 200
    assert overview.json()["documents"] == 0

    dashboard = client.get("/admin/")
    assert dashboard.status_code == 200
    assert "NorthStar Control Center" in dashboard.text
    assert client.get("/admin/assets/styles.css").status_code == 200
    dashboard_script = client.get("/admin/assets/app.js")
    assert dashboard_script.status_code == 200
    assert "WhatsApp messaging control" in dashboard_script.text
    assert 'confirmAction({\n      title: "Delete approved answer?"' in dashboard_script.text
    assert "Â" not in dashboard_script.text
    assert "â" not in dashboard_script.text
    for route in ("overview", "knowledge", "training", "feedback", "enrollments", "activity", "configuration"):
        assert f'data-route="{route}"' in dashboard.text

    public = client.get("/v1/public/config")
    assert public.status_code == 200
    feedback = public.json()["feedback"]
    assert feedback["enabled"] is True
    assert feedback["whatsapp_url"].startswith("https://wa.me/919876543210?text=FEEDBACK")

    repaired_prefill = client.put(
        "/v1/admin/configuration",
        headers=headers,
        json={"version": 0, "settings": {"whatsapp_feedback_prefill": "Please report an issue here"}},
    )
    assert repaired_prefill.status_code == 200
    assert repaired_prefill.json()["settings"]["whatsapp_feedback_prefill"] == (
        "FEEDBACK\nPlease report an issue here"
    )

    invalid_type = client.put(
        "/v1/admin/configuration",
        headers=headers,
        json={"version": repaired_prefill.json()["version"], "settings": {"top_k": None}},
    )
    assert invalid_type.status_code == 422


def test_approved_answer_lifecycle_and_whatsapp_preview(admin_client, monkeypatch):
    client, headers = admin_client
    created = client.post(
        "/v1/admin/approved-answers",
        headers=headers,
        json={"course": "CMA", "question": "What is variance analysis?", "answer": "It compares actual and planned results."},
    )
    assert created.status_code == 201, created.text
    record = created.json()
    assert record["status"] == "draft"

    published = client.post(f"/v1/admin/approved-answers/{record['id']}/publish", headers=headers)
    assert published.status_code == 200
    assert published.json()["status"] == "published"

    class FakeMentor:
        async def answer(self, request):
            match = await AdminStore().find_published_answer(request.course, request.message)
            return ChatResponse(answer=match["answer"], sources=[])

        async def aclose(self):
            return None

    monkeypatch.setattr(admin_api_module, "MentorService", FakeMentor)
    preview = client.post(
        "/v1/admin/mentor/whatsapp-preview",
        headers=headers,
        json={"course": "CMA", "mode": "doubt_solving", "level": "beginner", "question": "What is variance analysis?"},
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["exact_match"] is True
    assert "*CMA | Doubt Solving*" in preview.json()["whatsapp_answer"]
    assert "It compares actual and planned results." in preview.json()["whatsapp_answer"]

    deleted = client.delete(f"/v1/admin/approved-answers/{record['id']}", headers=headers)
    assert deleted.status_code == 204
    assert client.get("/v1/admin/approved-answers", headers=headers).json()["total"] == 0


def test_whatsapp_dashboard_master_switch_and_status(admin_client):
    client, headers = admin_client
    initial = client.get("/v1/admin/whatsapp/status", headers=headers)
    assert initial.status_code == 200
    status_payload = initial.json()
    assert status_payload["enabled"] is True
    assert status_payload["open_cma_access"] is True
    assert status_payload["outbound_ready"] is True
    assert status_payload["webhook_callback_url"] == "https://example.test/webhooks/whatsapp"
    assert status_payload["direct_callback_path"] == "/v1/whatsapp/ziplin/webhook"
    assert status_payload["relay_path"] == "/v1/whatsapp/ziplin/relay"
    assert status_payload["delivery_mode"] == "existing_webhook_relay"
    assert status_payload["routing_ready"] is True
    assert status_payload["stable_ingress_configured"] is True
    assert status_payload["relay_destination_url"] == (
        "https://mentor.example.test/v1/whatsapp/ziplin/relay"
    )
    assert "access_token" not in status_payload

    paused = client.put(
        "/v1/admin/whatsapp/messaging",
        headers=headers,
        json={"enabled": False, "version": status_payload["configuration_version"]},
    )
    assert paused.status_code == 200, paused.text
    assert paused.json()["enabled"] is False

    webhook_body = json.dumps({"object": "whatsapp_business_account", "entry": []}).encode()
    webhook_signature = "sha256=" + hmac.new(b"test-app-secret", webhook_body, hashlib.sha256).hexdigest()
    webhook = client.post(
        "/v1/whatsapp/webhook",
        content=webhook_body,
        headers={"content-type": "application/json", "x-hub-signature-256": webhook_signature},
    )
    assert webhook.status_code == 200
    assert webhook.json() == {"status": "paused", "event_id": None}
    assert client.get("/v1/admin/whatsapp/status", headers=headers).json()["queue"]["counts"]["pending"] == 0

    public = client.get("/v1/public/config")
    assert public.json()["feedback"]["enabled"] is False
    outbound = client.post(
        "/v1/admin/whatsapp/send-hi",
        headers=headers,
        json={"to": "919876543210"},
    )
    assert outbound.status_code == 409
    assert "paused" in outbound.json()["detail"].lower()

    resumed = client.put(
        "/v1/admin/whatsapp/messaging",
        headers=headers,
        json={"enabled": True, "version": paused.json()["configuration_version"]},
    )
    assert resumed.status_code == 200
    assert resumed.json()["enabled"] is True


def test_public_feedback_is_available_for_a_stable_authenticated_relay(admin_client, monkeypatch):
    client, _headers = admin_client
    monkeypatch.setenv("WHATSAPP_APP_SECRET", "")
    monkeypatch.setenv("META_APP_SECRET", "")
    monkeypatch.setenv("WHATSAPP_RELAY_TOKEN", "relay-secret-that-is-longer-than-32-characters")
    monkeypatch.setenv("NORTHSTAR_PUBLIC_BASE_URL", "https://mentor.example.test")
    get_settings.cache_clear()

    response = client.get("/v1/public/config")

    assert response.status_code == 200
    assert response.json()["feedback"]["enabled"] is True
    assert response.json()["feedback"]["whatsapp_url"].startswith("https://wa.me/919876543210")


def test_document_upload_edit_conflict_and_delete_api(admin_client):
    client, headers = admin_client
    upload = client.post(
        "/v1/admin/knowledge/documents",
        headers=headers,
        data={"course": "CMA", "doc_type": "lesson"},
        files={"files": ("variance.md", b"# Variance\nA useful lesson.", "text/markdown")},
    )
    assert upload.status_code == 201, upload.text
    document = upload.json()["items"][0]
    document_id = document["id"]

    listing = client.get("/v1/admin/knowledge/documents?course=CMA", headers=headers)
    assert listing.status_code == 200
    assert listing.json()["total"] == 1

    detail = client.get(f"/v1/admin/knowledge/documents/{document_id}", headers=headers).json()
    update_payload = {
        "title": "Variance analysis",
        "course": "CMA",
        "doc_type": "notes",
        "content": "Updated lesson content for retrieval.",
        "version": detail["version"],
    }
    updated = client.put(
        f"/v1/admin/knowledge/documents/{document_id}", headers=headers, json=update_payload
    )
    assert updated.status_code == 200
    assert updated.json()["version"] == 2
    assert updated.json()["status"] == "draft"

    stale = client.put(
        f"/v1/admin/knowledge/documents/{document_id}", headers=headers, json=update_payload
    )
    assert stale.status_code == 409

    deleted = client.delete(f"/v1/admin/knowledge/documents/{document_id}", headers=headers)
    assert deleted.status_code == 204
    assert client.get(f"/v1/admin/knowledge/documents/{document_id}", headers=headers).status_code == 404


def test_configuration_is_versioned_atomic_and_audited(admin_client):
    client, headers = admin_client
    current = client.get("/v1/admin/configuration", headers=headers)
    assert current.status_code == 200
    original = current.json()

    saved = client.put(
        "/v1/admin/configuration",
        headers=headers,
        json={
            "version": original["version"],
            "settings": {"top_k": 7},
            "system_prompt": original["system_prompt"],
        },
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["version"] == original["version"] + 1
    assert saved.json()["settings"]["top_k"] == 7

    stale = client.put(
        "/v1/admin/configuration",
        headers=headers,
        json={
            "version": original["version"],
            "settings": {"top_k": 3},
            "system_prompt": original["system_prompt"],
        },
    )
    assert stale.status_code == 409
    assert client.get("/v1/admin/configuration", headers=headers).json()["settings"]["top_k"] == 7

    activity = client.get("/v1/admin/audit?search=configuration.updated", headers=headers)
    assert activity.status_code == 200
    assert activity.json()["total"] == 1
    assert activity.json()["items"][0]["details"]["fields"] == ["top_k"]


def test_training_job_detail_and_retry_api(admin_client, monkeypatch):
    client, headers = admin_client
    scheduled: list[str] = []
    monkeypatch.setattr(admin_api_module, "schedule_training_job", scheduled.append)
    upload = client.post(
        "/v1/admin/knowledge/documents",
        headers=headers,
        data={"course": "CMA", "doc_type": "lesson"},
        files={"files": ("retry.md", b"Retryable indexing content.", "text/markdown")},
    )
    document_id = upload.json()["items"][0]["id"]
    queued = client.post(
        "/v1/admin/training/jobs",
        headers=headers,
        json={"document_ids": [document_id]},
    )
    assert queued.status_code == 202
    job_id = queued.json()["id"]
    assert scheduled == [job_id]

    with AdminStore()._connect() as connection:
        connection.execute(
            "UPDATE training_jobs SET status = 'failed', failed_documents = 1 WHERE id = ?",
            (job_id,),
        )
        connection.execute(
            "UPDATE knowledge_documents SET status = 'failed' WHERE id = ?",
            (document_id,),
        )

    detail = client.get(f"/v1/admin/training/jobs/{job_id}", headers=headers)
    assert detail.status_code == 200
    assert detail.json()["status"] == "failed"

    retried = client.post(f"/v1/admin/training/jobs/{job_id}/retry", headers=headers)
    assert retried.status_code == 202, retried.text
    assert retried.json()["id"] != job_id
    assert retried.json()["document_ids"] == [document_id]
    assert scheduled[-1] == retried.json()["id"]


def test_feedback_attachment_requires_admin_auth(admin_client):
    client, headers = admin_client
    feedback = asyncio.run(
        AdminStore().create_feedback(
            whatsapp_message_id="wamid.api-feedback",
            sender="919999999999",
            profile_name="Student",
            message="There is an error in this answer.",
            image_bytes=b"\x89PNG\r\n\x1a\nvalid-test-image",
            mime_type="image/png",
        )
    )
    path = f"/v1/admin/feedback/{feedback['id']}/attachment"

    assert client.get(path).status_code == 401
    attachment = client.get(path, headers=headers)
    assert attachment.status_code == 200
    assert attachment.headers["content-type"].startswith("image/png")
    assert attachment.content.startswith(b"\x89PNG")


def test_admin_enrollment_api_writes_and_revokes_in_excel(admin_client, tmp_path, monkeypatch):
    client, headers = admin_client
    workbook_path = tmp_path / "whatsapp_enrollments.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Enrollments"
    sheet.append(["phone_number", "course", "active", "student_name", "notes"])
    sheet.append(["919916039894", "CMA", "YES", "Student", "Keep this note"])
    workbook.save(workbook_path)
    workbook.close()
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_FILE", str(workbook_path))
    get_settings.cache_clear()
    whatsapp_module._load_enrollment_workbook.cache_clear()

    granted = client.put(
        "/v1/admin/whatsapp/enrollments/919916039894",
        headers=headers,
        json={"course": "CPA"},
    )
    assert granted.status_code == 200, granted.text
    assert granted.json()["source"] == "excel"
    assert granted.json()["courses"] == ["CMA", "CPA"]

    revoked = client.delete(
        "/v1/admin/whatsapp/enrollments/919916039894/CMA",
        headers=headers,
    )
    assert revoked.status_code == 200, revoked.text
    assert revoked.json()["source"] == "excel"
    assert revoked.json()["courses"] == ["CPA"]

    saved = load_workbook(workbook_path, read_only=True, data_only=True)
    rows = list(saved["Enrollments"].iter_rows(min_row=2, values_only=True))
    saved.close()
    states = {(str(row[0]), str(row[1])): str(row[2]) for row in rows}
    assert states[("919916039894", "CMA")] == "NO"
    assert states[("919916039894", "CPA")] == "YES"
    assert rows[0][4] == "Keep this note"
    activity = client.get(
        "/v1/admin/audit?search=919916039894",
        headers=headers,
    )
    assert activity.status_code == 200
    assert [item["action"] for item in activity.json()["items"]] == [
        "enrollment.revoked",
        "enrollment.granted",
    ]


def test_admin_enrollment_lookup_surfaces_corrupt_excel(admin_client, tmp_path, monkeypatch):
    client, headers = admin_client
    workbook_path = tmp_path / "whatsapp_enrollments.xlsx"
    workbook_path.write_bytes(b"not-an-xlsx-workbook")
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_FILE", str(workbook_path))
    get_settings.cache_clear()
    whatsapp_module._load_enrollment_workbook.cache_clear()

    response = client.get(
        "/v1/admin/whatsapp/enrollments/919916039894",
        headers=headers,
    )

    assert response.status_code == 409
    assert "could not be read" in response.json()["detail"]


def test_admin_enrollment_endpoints_reject_invalid_e164_length(admin_client):
    client, headers = admin_client

    too_short = client.get("/v1/admin/whatsapp/enrollments/123", headers=headers)
    too_long = client.get(
        "/v1/admin/whatsapp/enrollments/1234567890123456",
        headers=headers,
    )

    assert too_short.status_code == 422
    assert too_long.status_code == 422


def test_production_startup_rejects_default_or_missing_secrets(monkeypatch, tmp_path):
    monkeypatch.setenv("APP_ENVIRONMENT", "production")
    monkeypatch.setenv("RUNTIME_CONFIG_FILE", str(tmp_path / "runtime-config.json"))
    monkeypatch.setenv("ADMIN_TOKEN", "change-me")
    monkeypatch.setenv("WHATSAPP_VERIFY_TOKEN", "change-me-whatsapp")
    monkeypatch.setenv("WHATSAPP_ACCESS_TOKEN", "")
    monkeypatch.setenv("WHATSAPP_TOKEN", "")
    monkeypatch.setenv("WHATSAPP_PHONE_NUMBER_ID", "")
    monkeypatch.setenv("WHATSAPP_APP_SECRET", "")
    monkeypatch.setenv("META_APP_SECRET", "")
    monkeypatch.setenv("WHATSAPP_FEEDBACK_NUMBER", "not-a-number")
    monkeypatch.setenv("WHATSAPP_FEEDBACK_PREFILL", "Please report a problem")
    monkeypatch.setenv("NVIDIA_API_KEY", "")
    monkeypatch.setenv("GEMINI_API_KEY", "")
    monkeypatch.setenv("REQUEST_TIMEOUT_SECONDS", "120")
    monkeypatch.setenv("TRAINING_LEASE_TIMEOUT_SECONDS", "300")
    monkeypatch.setenv("WHATSAPP_WEBHOOK_LEASE_TIMEOUT_SECONDS", "10")
    monkeypatch.setenv("WHATSAPP_WEBHOOK_RECOVERY_INTERVAL_SECONDS", "15")
    monkeypatch.setenv("WHATSAPP_WEBHOOK_MAX_ATTEMPTS", "0")
    get_settings.cache_clear()

    with pytest.raises(RuntimeError, match="Unsafe production configuration") as error:
        validate_startup_configuration()
    assert "E.164" in str(error.value)
    assert "FEEDBACK command marker" in str(error.value)
    assert "GEMINI_API_KEY is required for WhatsApp image questions" in str(error.value)
    assert "TRAINING_LEASE_TIMEOUT_SECONDS must be at least 480" in str(error.value)
    assert "WHATSAPP_WEBHOOK_LEASE_TIMEOUT_SECONDS" in str(error.value)
    assert "WHATSAPP_WEBHOOK_MAX_ATTEMPTS" in str(error.value)
    get_settings.cache_clear()


def test_readiness_probes_redis_and_qdrant(admin_client, monkeypatch):
    client, _ = admin_client
    checked = {"redis_closed": False, "qdrant_closed": False, "qdrant_schema": False}

    class FakeRedis:
        async def ping(self):
            return True

        async def aclose(self):
            checked["redis_closed"] = True

    class FakeCache:
        def __init__(self):
            self.redis = FakeRedis()

    class FakeQdrantClient:
        async def get_collections(self):
            return object()

        async def close(self):
            checked["qdrant_closed"] = True

    class FakeVectorStore:
        def __init__(self):
            self.client = FakeQdrantClient()

        async def ensure_collection(self):
            checked["qdrant_schema"] = True

    monkeypatch.setattr(main_module, "Cache", FakeCache)
    monkeypatch.setattr(main_module, "VectorStore", FakeVectorStore)
    response = client.get("/ready")
    assert response.status_code == 200
    assert response.json()["status"] == "ready"
    assert response.json()["dependencies"]["sqlite"] == "ok"
    assert checked == {"redis_closed": True, "qdrant_closed": True, "qdrant_schema": True}


def test_whatsapp_webhook_is_durably_enqueued_before_acceptance(admin_client, monkeypatch):
    client, _ = admin_client
    scheduled = []
    monkeypatch.setattr(main_module, "schedule_whatsapp_webhook", scheduled.append)
    payload = {"entry": [{"id": "entry-api", "changes": []}]}
    body = json.dumps(payload, separators=(",", ":")).encode("utf-8")
    signature = "sha256=" + hmac.new(
        b"test-app-secret",
        body,
        hashlib.sha256,
    ).hexdigest()

    response = client.post(
        "/v1/whatsapp/webhook",
        content=body,
        headers={
            "content-type": "application/json",
            "x-hub-signature-256": signature,
        },
    )
    assert response.status_code == 200
    event_id = response.json()["event_id"]
    assert scheduled == [event_id]
    with AdminStore()._connect() as connection:
        row = connection.execute(
            "SELECT status, payload FROM whatsapp_webhook_events WHERE id = ?",
            (event_id,),
        ).fetchone()
    assert row["status"] == "pending"
    assert "entry-api" in row["payload"]


def test_existing_webhook_can_relay_into_the_same_durable_queue(admin_client, monkeypatch):
    client, _ = admin_client
    scheduled = []
    monkeypatch.setenv("WHATSAPP_RELAY_TOKEN", "relay-secret-that-is-longer-than-32-characters")
    get_settings.cache_clear()
    monkeypatch.setattr(main_module, "schedule_whatsapp_webhook", scheduled.append)
    payload = {"entry": [{"id": "entry-from-existing-webhook", "changes": []}]}

    unauthorized = client.post("/v1/whatsapp/relay", json=payload)
    assert unauthorized.status_code == 401

    response = client.post(
        "/v1/whatsapp/relay",
        json=payload,
        headers={"x-northstar-relay-token": "relay-secret-that-is-longer-than-32-characters"},
    )
    assert response.status_code == 200
    event_id = response.json()["event_id"]
    assert scheduled == [event_id]
    with AdminStore()._connect() as connection:
        row = connection.execute(
            "SELECT status, payload FROM whatsapp_webhook_events WHERE id = ?",
            (event_id,),
        ).fetchone()
    assert row["status"] == "pending"
    assert "entry-from-existing-webhook" in row["payload"]


def test_ziplin_relay_status_requires_a_stable_public_destination(admin_client, monkeypatch):
    client, headers = admin_client
    monkeypatch.setenv("WHATSAPP_RELAY_TOKEN", "relay-secret-that-is-longer-than-32-characters")
    monkeypatch.setenv("NORTHSTAR_PUBLIC_BASE_URL", "https://mentor.example.test/")
    get_settings.cache_clear()

    payload = client.get("/v1/admin/whatsapp/status", headers=headers).json()

    assert payload["delivery_mode"] == "existing_webhook_relay"
    assert payload["stable_ingress_configured"] is True
    assert payload["routing_ready"] is True
    assert payload["relay_destination_url"] == (
        "https://mentor.example.test/v1/whatsapp/ziplin/relay"
    )

    monkeypatch.setenv(
        "NORTHSTAR_PUBLIC_BASE_URL",
        "https://temporary-example.trycloudflare.com",
    )
    get_settings.cache_clear()
    temporary = client.get("/v1/admin/whatsapp/status", headers=headers).json()
    assert temporary["stable_ingress_configured"] is False
    assert temporary["routing_ready"] is False
    assert temporary["relay_destination_url"] == ""


def test_ziplin_relay_accepts_only_the_configured_phone_number(admin_client, monkeypatch):
    client, _headers = admin_client
    relay_token = "ziplin-relay-secret-that-is-longer-than-32-characters"
    monkeypatch.setenv("WHATSAPP_RELAY_TOKEN", relay_token)
    get_settings.cache_clear()
    scheduled: list[str] = []
    monkeypatch.setattr(main_module, "schedule_whatsapp_webhook", scheduled.append)

    def payload(phone_id: str, message_id: str) -> dict:
        return {
            "object": "whatsapp_business_account",
            "entry": [
                {
                    "id": "test-waba",
                    "changes": [
                        {
                            "field": "messages",
                            "value": {
                                "metadata": {"phone_number_id": phone_id},
                                "contacts": [{"wa_id": "919535210826"}],
                                "messages": [
                                    {
                                        "id": message_id,
                                        "from": "919535210826",
                                        "timestamp": "1700000000",
                                        "type": "text",
                                        "text": {"body": "Hi"},
                                    }
                                ],
                            },
                        }
                    ],
                }
            ],
        }

    wrong = client.post(
        "/v1/whatsapp/ziplin/relay",
        json=payload("other-northstar-phone-id", "wrong-number-message"),
        headers={"x-ziplin-relay-token": relay_token},
    )
    assert wrong.status_code == 200
    assert wrong.json() == {
        "status": "ignored",
        "event_id": None,
        "reason": "different_phone_number",
    }
    assert scheduled == []

    ziplin = client.post(
        "/v1/whatsapp/ziplin/relay",
        json=payload("test-phone-id", "ziplin-message"),
        headers={"x-ziplin-relay-token": relay_token},
    )
    assert ziplin.status_code == 200
    assert ziplin.json()["status"] == "accepted"
    assert scheduled == [ziplin.json()["event_id"]]


def test_ziplin_direct_webhook_uses_dedicated_alias(admin_client):
    client, _headers = admin_client
    verification = client.get(
        "/v1/whatsapp/ziplin/webhook",
        params={
            "hub.mode": "subscribe",
            "hub.verify_token": "test-verify-token",
            "hub.challenge": "12345",
        },
    )
    # The fixture does not set a verify token, so the alias must reject the
    # request through the same verification path as the legacy endpoint.
    assert verification.status_code == 403


def test_relay_is_closed_when_shared_secret_is_not_configured(admin_client, monkeypatch):
    client, _ = admin_client
    monkeypatch.setenv("WHATSAPP_RELAY_TOKEN", "")
    get_settings.cache_clear()
    response = client.post("/v1/whatsapp/relay", json={"entry": []})
    assert response.status_code == 503


def test_course_helper_preserves_full_international_number_and_saves_workbook(tmp_path):
    path = tmp_path / "enrollments.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Enrollments"
    sheet.append(["phone_number", "course", "active", "student_name", "notes"])
    workbook.create_sheet("Instructions")
    workbook.save(path)
    set_student_courses(path, "+44 7700 900123", ["CMA", "CPA"], "Test student")

    saved = load_workbook(path, read_only=True)
    try:
        rows = list(saved["Enrollments"].iter_rows(min_row=2, values_only=True))
    finally:
        saved.close()
    assert normalize_phone("+44 7700 900123") == "447700900123"
    assert {(row[0], row[1], row[2]) for row in rows} == {
        ("447700900123", "CMA", "YES"),
        ("447700900123", "CPA", "YES"),
    }
    with pytest.raises(ValueError, match="8 to 15 digits"):
        normalize_phone("123")


def test_usage_analytics_groups_students_courses_channels_and_latency(admin_client):
    client, headers = admin_client
    store = AdminStore()
    asyncio.run(store.record_usage(student_id="student-a", channel="web", course="CMA", mode="teach", status="success", latency_ms=1200))
    asyncio.run(store.record_usage(student_id="student-a", channel="whatsapp", course="CMA", mode="quiz", status="error", latency_ms=1800))
    asyncio.run(store.record_usage(student_id="student-b", channel="web", course="CPA", mode="revise", status="success", latency_ms=900))

    response = client.get("/v1/admin/analytics?days=30", headers=headers)
    assert response.status_code == 200
    payload = response.json()
    assert payload["totals"] == {
        "requests": 3,
        "students": 2,
        "errors": 1,
        "success_rate": 66.7,
        "average_latency_ms": 1300,
    }
    assert {item["label"]: item["value"] for item in payload["courses"]} == {"CMA": 2, "CPA": 1}
    assert payload["students"][0]["student_id"] == "student-a"


def test_admin_can_run_live_model_health_contract(admin_client, monkeypatch):
    client, headers = admin_client

    class FakeMentorLLM:
        async def answer_text(self, **_kwargs):
            return "Variance analysis highlights differences between planned and actual performance."

        async def aclose(self):
            return None

    monkeypatch.setattr(admin_api_module, "MentorLLMService", FakeMentorLLM)
    response = client.post("/v1/admin/model/test", headers=headers, json={})
    assert response.status_code == 200
    assert response.json()["status"] == "ok"
    assert response.json()["provider"] == "nvidia"
    assert "Variance analysis" in response.json()["answer"]


def test_bulk_excel_enrollment_import_updates_authority_and_roster(admin_client, monkeypatch, tmp_path):
    client, headers = admin_client
    authority = tmp_path / "authority.xlsx"
    base = Workbook()
    base.active.title = "Enrollments"
    base.active.append(["phone_number", "course", "active", "student_name", "notes"])
    base.save(authority)
    monkeypatch.setenv("WHATSAPP_ENROLLMENTS_FILE", str(authority))
    get_settings.cache_clear()
    whatsapp_module._load_enrollment_workbook.cache_clear()

    upload = Workbook()
    upload.active.title = "Enrollments"
    upload.active.append(["phone_number", "course", "active", "student_name", "notes"])
    upload.active.append(["447700900123", "CMA", "YES", "Student One", "Imported"])
    upload.active.append(["447700900123", "CPA", "YES", "Student One", "Imported"])
    from io import BytesIO
    body = BytesIO()
    upload.save(body)

    response = client.post(
        "/v1/admin/whatsapp/enrollments/import",
        headers=headers,
        files={"file": ("enrollments.xlsx", body.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200
    assert response.json()["rows"] == 2
    roster = client.get("/v1/admin/whatsapp/enrollments", headers=headers)
    assert roster.status_code == 200
    assert roster.json()["items"] == [{"phone": "447700900123", "courses": ["CMA", "CPA"], "source": "excel"}]
